import unittest
import os
import json
import sqlite3
import tempfile
import io
from werkzeug.security import generate_password_hash

import app

class RouteAndSecurityDeepScan(unittest.TestCase):
    def setUp(self):
        # Create a fresh temp database for testing
        self.db_fd, self.db_path = tempfile.mkstemp(suffix=".db")
        app.DATABASE = self.db_path
        app.app.config['TESTING'] = True
        app.app.config['SECRET_KEY'] = 'test-secret-key-12345'
        app.app.config['WTF_CSRF_ENABLED'] = False
        self.upload_dir = tempfile.mkdtemp()
        app.app.config['UPLOAD_FOLDER'] = self.upload_dir
        
        with app.app.app_context():
            app.init_db()
            db = app.get_db()
            # Setup specific test accounts
            # Student 1 (Verified)
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (10, 'Student One', 's1@test.com', ?, 'student', 'CS', '6', ?, 1, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            # Student 2 (Verified)
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (11, 'Student Two', 's2@test.com', ?, 'student', 'CS', '6', ?, 1, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            # Student 3 (Unverified)
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (12, 'Student Unverified', 's3@test.com', ?, 'student', 'CS', '6', ?, 0, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            # Staff 1
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (20, 'Staff Asha', 'staff1@test.com', ?, 'staff', 'CS', 'All', ?, 1, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            # Staff 2 (New Staff with 0 notes)
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (21, 'Staff New', 'staff2@test.com', ?, 'staff', 'Math', 'All', ?, 1, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            # Admin
            db.execute("INSERT OR REPLACE INTO users (id, name, email, password_hash, role, department, semester, last_login, is_verified, graduation_status) VALUES (30, 'Admin Super', 'admin1@test.com', ?, 'admin', 'Admin', 'All', ?, 1, 'active')", (generate_password_hash('pass12345'), app.now_iso()))
            db.commit()

        self.client = app.app.test_client()

    def tearDown(self):
        os.close(self.db_fd)
        if os.path.exists(self.db_path):
            try: os.remove(self.db_path)
            except: pass
        if os.path.exists(self.upload_dir):
            try:
                import shutil
                shutil.rmtree(self.upload_dir)
            except: pass

    def login_as(self, user_id):
        with self.client.session_transaction() as sess:
            sess['user_id'] = user_id
            sess['_csrf_token'] = 'test-csrf-token'

    def get_csrf(self):
        return 'test-csrf-token'

    def test_unverified_user_cannot_login(self):
        with self.client:
            res = self.client.post('/login', data={'email': 's3@test.com', 'password': 'pass12345', '_csrf_token': self.get_csrf()}, follow_redirects=True)
            self.assertIn(b"waiting for administrator verification", res.data)

    def test_rbac_restrictions(self):
        # Student cannot access admin
        self.login_as(10)
        res = self.client.get('/admin')
        self.assertEqual(res.status_code, 403)
        
        # Student cannot access staff analytics
        res = self.client.get('/staff/analytics')
        self.assertEqual(res.status_code, 403)

        # Staff can access analytics
        self.login_as(20)
        res = self.client.get('/staff/analytics')
        self.assertEqual(res.status_code, 200)

        # Staff cannot access admin
        res = self.client.get('/admin')
        self.assertEqual(res.status_code, 403)

        # Admin can access admin
        self.login_as(30)
        res = self.client.get('/admin')
        self.assertEqual(res.status_code, 200)

    def test_staff_dashboard_with_zero_documents(self):
        # Staff 2 has no documents, math dept has no documents, public has no documents
        self.login_as(21)
        try:
            res = self.client.get('/dashboard')
            self.assertEqual(res.status_code, 200)
        except Exception as e:
            self.fail(f"Staff dashboard raised exception with 0 documents: {e}")

    def test_staff_export_roster_with_zero_documents(self):
        self.login_as(21)
        try:
            res = self.client.get('/staff/export-roster')
            self.assertEqual(res.status_code, 200)
            self.assertEqual(res.mimetype, 'text/csv')
        except Exception as e:
            self.fail(f"Export roster raised exception with 0 documents: {e}")

    def test_document_crud_and_access_modes(self):
        # Student 1 creates public document
        self.login_as(10)
        res = self.client.post('/documents/new', data={
            'title': 'Operating Systems Intro',
            'subject': 'Operating Systems',
            'department': 'CS',
            'semester': '6',
            'content': 'Processes and threads are fundamental execution units in an operating system.',
            'visibility': 'public',
            'access_mode': 'all',
            '_csrf_token': self.get_csrf()
        }, follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        with app.app.app_context():
            doc = app.get_db().execute("SELECT id FROM documents WHERE title = 'Operating Systems Intro'").fetchone()
            self.assertIsNotNone(doc)
            doc_id = doc['id']

        # Student 2 can view public document
        self.login_as(11)
        res = self.client.get(f'/documents/{doc_id}')
        self.assertEqual(res.status_code, 200)

        # Student 1 creates private document with allow_list for Student 2
        self.login_as(10)
        res = self.client.post('/documents/new', data={
            'title': 'Private Algorithms Research',
            'subject': 'Algorithms',
            'department': 'CS',
            'semester': '6',
            'content': 'Dynamic Programming optimizes overlapping subproblems via memoization.',
            'visibility': 'private',
            'access_mode': 'allow_list',
            'student_ids': ['11'],
            '_csrf_token': self.get_csrf()
        }, follow_redirects=True)
        self.assertEqual(res.status_code, 200)

        with app.app.app_context():
            priv_doc = app.get_db().execute("SELECT id FROM documents WHERE title = 'Private Algorithms Research'").fetchone()
            priv_doc_id = priv_doc['id']

        # Student 2 is on allow_list -> can view
        self.login_as(11)
        res = self.client.get(f'/documents/{priv_doc_id}')
        self.assertEqual(res.status_code, 200)

        # What about Staff or Admin viewing private/allow-listed document?
        self.login_as(30) # Admin
        res = self.client.get(f'/documents/{priv_doc_id}')
        # Check admin access behavior
        print("Admin access to private allow-listed doc status:", res.status_code)

    def test_viva_evaluation_route(self):
        self.login_as(10)
        # Create a document
        with app.app.app_context():
            db = app.get_db()
            cursor = db.execute("INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('DBMS', 'CS', 'CS', '6', 'Normalization reduces data redundancy in relational databases.', 'public', 0, 10, ?, 'ready', 'all')", (app.now_iso(),))
            doc_id = cursor.lastrowid
            db.commit()

        res = self.client.post(f'/documents/{doc_id}/viva/evaluate', data={
            'question_id': '1',
            'student_answer': 'Normalization minimizes data redundancy and improves database integrity.',
            '_csrf_token': self.get_csrf()
        })
        self.assertEqual(res.status_code, 200)
        data = json.loads(res.data)
        self.assertTrue(data.get('success'))
        self.assertIn('evaluation', data)

    def test_flashcard_review_route(self):
        self.login_as(10)
        with app.app.app_context():
            db = app.get_db()
            cursor = db.execute("INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('DBMS', 'CS', 'CS', '6', 'Normalization reduces data redundancy in relational databases.', 'public', 0, 10, ?, 'ready', 'all')", (app.now_iso(),))
            doc_id = cursor.lastrowid
            db.commit()

        res = self.client.post(f'/documents/{doc_id}/flashcards/review', data={
            'card_id': 'card123',
            'rating': 'good',
            '_csrf_token': self.get_csrf()
        })
        self.assertEqual(res.status_code, 200)
        data = json.loads(res.data)
        self.assertTrue(data.get('success'))
        self.assertEqual(data.get('new_box'), 2)
    def test_delete_quiz_and_regeneration(self):
        self.login_as(10)
        with app.app.app_context():
            db = app.get_db()
            cursor = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('DBMS', 'CS', 'CS', '6', 'Normalization reduces data redundancy in relational databases. SQL is structured query language for data manipulation.', 'public', 0, 10, ?, 'ready', 'all')",
                (app.now_iso(),)
            )
            doc_id = cursor.lastrowid
            db.execute("INSERT INTO study_summaries (document_id, position, sentence) VALUES (?, 0, 'Key summary sentence.')", (doc_id,))
            db.execute("INSERT INTO quiz_questions (document_id, position, prompt, options_json, answer) VALUES (?, 0, 'What is SQL?', '[\"Structured Query Language\", \"Simple Query Link\"]', 'Structured Query Language')", (doc_id,))
            db.execute("INSERT INTO quiz_attempts (document_id, user_id, score, total, created_at) VALUES (?, 10, 1, 1, ?)", (doc_id, app.now_iso()))
            db.commit()

        res = self.client.get(f'/documents/{doc_id}/quiz')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b'Delete Quiz Bank', res.data)
        self.assertIn(b'What is SQL?', res.data)

        del_res = self.client.post(f'/documents/{doc_id}/quiz/delete', headers={'Referer': f'/documents/{doc_id}/quiz'}, data={'_csrf_token': self.get_csrf()})
        self.assertEqual(del_res.status_code, 302)
        self.assertIn(f'/documents/{doc_id}/quiz', del_res.headers['Location'])

        with app.app.app_context():
            db = app.get_db()
            q_count = db.execute("SELECT COUNT(*) FROM quiz_questions WHERE document_id = ?", (doc_id,)).fetchone()[0]
            a_count = db.execute("SELECT COUNT(*) FROM quiz_attempts WHERE document_id = ?", (doc_id,)).fetchone()[0]
            s_count = db.execute("SELECT COUNT(*) FROM study_summaries WHERE document_id = ?", (doc_id,)).fetchone()[0]
            self.assertEqual(q_count, 0)
            self.assertEqual(a_count, 0)
            self.assertEqual(s_count, 1)

        quiz_res = self.client.get(f'/documents/{doc_id}/quiz')
        self.assertEqual(quiz_res.status_code, 200)
        self.assertIn(b'Question bank is ready for generation', quiz_res.data)
        self.assertIn(b'Generate study pack', quiz_res.data)
        self.assertNotIn(b'Delete Quiz Bank', quiz_res.data)

        with app.app.app_context():
            db = app.get_db()
            q_count = db.execute("SELECT COUNT(*) FROM quiz_questions WHERE document_id = ?", (doc_id,)).fetchone()[0]
            self.assertEqual(q_count, 0)

        gen_res = self.client.post(f'/documents/{doc_id}/study-pack/generate', headers={'Referer': f'/documents/{doc_id}/quiz'}, data={'_csrf_token': self.get_csrf()})
        self.assertEqual(gen_res.status_code, 302)

        with app.app.app_context():
            db = app.get_db()
            q_count = db.execute("SELECT COUNT(*) FROM quiz_questions WHERE document_id = ?", (doc_id,)).fetchone()[0]
            self.assertGreater(q_count, 0)

    def test_delete_quiz_unauthorized(self):
        self.login_as(11)
        with app.app.app_context():
            db = app.get_db()
            cursor = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('DBMS', 'CS', 'CS', '6', 'Normalization reduces data redundancy.', 'public', 0, 10, ?, 'ready', 'all')",
                (app.now_iso(),)
            )
            doc_id = cursor.lastrowid
            db.execute("INSERT INTO quiz_questions (document_id, position, prompt, options_json, answer) VALUES (?, 0, 'What is SQL?', '[\"Structured Query Language\"]', 'Structured Query Language')", (doc_id,))
            db.commit()

        del_res = self.client.post(f'/documents/{doc_id}/quiz/delete', data={'_csrf_token': self.get_csrf()})
        self.assertEqual(del_res.status_code, 403)

    def test_staff_multi_format_csv_exports(self):
        self.login_as(20)  # Staff 1
        with app.app.app_context():
            db = app.get_db()
            cur = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('Database Systems', 'DBMS', 'CS', '6', 'Relational database management principles.', 'public', 0, 20, ?, 'ready', 'all')",
                (app.now_iso(),)
            )
            doc_id = cur.lastrowid
            # Record views and attempts by student 10
            db.execute("INSERT INTO activity_logs (user_id, document_id, event_type, detail, created_at) VALUES (10, ?, 'document_view', 'viewed', ?)", (doc_id, app.now_iso()))
            db.execute("INSERT INTO quiz_attempts (document_id, user_id, score, total, created_at) VALUES (?, 10, 9, 10, ?)", (doc_id, app.now_iso()))
            db.commit()

        # 1. Summary Gradebook CSV
        res_summary = self.client.get('/staff/export-roster')
        self.assertEqual(res_summary.status_code, 200)
        self.assertEqual(res_summary.mimetype, 'text/csv')
        # Check UTF-8 BOM
        self.assertTrue(res_summary.data.startswith(b'\xef\xbb\xbf'))
        csv_text = res_summary.data.decode('utf-8')
        self.assertIn("Student ID", csv_text)
        self.assertIn("Student One", csv_text)
        self.assertIn("Distinction (85-100%)", csv_text)

        # 2. Materials Audit CSV
        res_materials = self.client.get('/staff/export-roster?type=materials')
        self.assertEqual(res_materials.status_code, 200)
        self.assertEqual(res_materials.mimetype, 'text/csv')
        self.assertTrue(res_materials.data.startswith(b'\xef\xbb\xbf'))
        mat_text = res_materials.data.decode('utf-8')
        self.assertIn("Material ID", mat_text)
        self.assertIn("Database Systems", mat_text)
        self.assertIn("Exemplary", mat_text)

        # 3. Detailed Attempts Ledger CSV
        res_attempts = self.client.get('/staff/export/attempts')
        self.assertEqual(res_attempts.status_code, 200)
        self.assertEqual(res_attempts.mimetype, 'text/csv')
        self.assertTrue(res_attempts.data.startswith(b'\xef\xbb\xbf'))
        att_text = res_attempts.data.decode('utf-8')
        self.assertIn("Attempt ID", att_text)
        self.assertIn("Score Earned", att_text)
        self.assertIn("Student One", att_text)

    def test_staff_executive_report_routes(self):
        self.login_as(20)  # Staff 1
        with app.app.app_context():
            db = app.get_db()
            cur = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('Algorithms', 'CS', 'CS', '6', 'Sorting and search algorithms.', 'public', 0, 20, ?, 'ready', 'all')",
                (app.now_iso(),)
            )
            doc_id = cur.lastrowid
            db.execute("INSERT INTO quiz_attempts (document_id, user_id, score, total, created_at) VALUES (?, 10, 8, 10, ?)", (doc_id, app.now_iso()))
            db.commit()

        # View Executive Report in browser
        res_report = self.client.get('/staff/report')
        self.assertEqual(res_report.status_code, 200)
        self.assertIn(b"Executive Teaching Assessment Report", res_report.data)
        self.assertIn(b"Cohort Mastery & Grade Distribution", res_report.data)
        self.assertIn(b"Official Pedagogical Certification", res_report.data)
        self.assertIn(b"Print / Save as PDF", res_report.data)

        # Download Standalone HTML Report
        res_download = self.client.get('/staff/report/download')
        self.assertEqual(res_download.status_code, 200)
        self.assertEqual(res_download.mimetype, 'text/html')
        self.assertIn('attachment;', res_download.headers.get('Content-Disposition', ''))

    def test_staff_excel_exports(self):
        self.login_as(20)  # Staff 1
        with app.app.app_context():
            db = app.get_db()
            cur = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, created_at, processing_status, access_mode) VALUES ('Algorithms & DS', 'CS', 'CS', '6', 'Graph theory and sorting.', 'public', 0, 20, ?, 'ready', 'all')",
                (app.now_iso(),)
            )
            doc_id = cur.lastrowid
            db.execute("INSERT INTO quiz_attempts (document_id, user_id, score, total, created_at) VALUES (?, 10, 9, 10, ?)", (doc_id, app.now_iso()))
            db.commit()

        # 1. Master Excel Dossier (.xlsx)
        res_master = self.client.get('/staff/export-excel')
        self.assertEqual(res_master.status_code, 200)
        self.assertEqual(res_master.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', res_master.headers.get('Content-Disposition', ''))

        # Verify valid openpyxl workbook structure with all 3 sheets
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res_master.data))
        self.assertIn('Student Gradebook', wb.sheetnames)
        self.assertIn('Material Outcomes', wb.sheetnames)
        self.assertIn('Quiz Attempts Ledger', wb.sheetnames)

        # 2. Individual Gradebook Excel (.xlsx)
        res_gradebook = self.client.get('/staff/export/excel/gradebook')
        self.assertEqual(res_gradebook.status_code, 200)
        self.assertEqual(res_gradebook.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('gradebook', res_gradebook.headers.get('Content-Disposition', ''))

        # 3. Individual Materials Audit Excel (.xlsx)
        res_mat = self.client.get('/staff/export/excel/materials')
        self.assertEqual(res_mat.status_code, 200)
        self.assertEqual(res_mat.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('materials_audit', res_mat.headers.get('Content-Disposition', ''))

        # 4. Individual Detailed Attempts Excel (.xlsx)
        res_att = self.client.get('/staff/export/excel/attempts')
        self.assertEqual(res_att.status_code, 200)
        self.assertEqual(res_att.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('detailed_attempts', res_att.headers.get('Content-Disposition', ''))

    def test_staff_reports_role_security(self):
        # Student cannot access staff report or exports
        self.login_as(10)  # Student
        self.assertEqual(self.client.get('/staff/report').status_code, 403)
        self.assertEqual(self.client.get('/staff/report/download').status_code, 403)
        self.assertEqual(self.client.get('/staff/export-roster').status_code, 403)
        self.assertEqual(self.client.get('/staff/export-excel').status_code, 403)
        self.assertEqual(self.client.get('/staff/export/excel/gradebook').status_code, 403)

        # Admin can access staff report and exports
        self.login_as(30)  # Admin
        self.assertEqual(self.client.get('/staff/report').status_code, 200)
        self.assertEqual(self.client.get('/staff/report/download').status_code, 200)
        self.assertEqual(self.client.get('/staff/export-roster').status_code, 200)
        self.assertEqual(self.client.get('/staff/export-excel').status_code, 200)
        self.assertEqual(self.client.get('/staff/export/excel/gradebook').status_code, 200)


if __name__ == '__main__':
    unittest.main()

