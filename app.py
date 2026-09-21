import atexit
import csv
import difflib
import hashlib
import hmac
import io
import json
import logging
import os
import random
import re
import sqlite3
import secrets
import uuid
from urllib import error as urlerror
from urllib import request as urlrequest
from collections import Counter
from datetime import UTC, datetime, timedelta
from functools import wraps

from apscheduler.schedulers.background import BackgroundScheduler
from flask import Flask, Response, abort, flash, g, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

import zipfile
import xml.etree.ElementTree as ET


BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DATABASE = os.path.join(BASE_DIR, "omnistudy.db")
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
ALLOWED_EXTENSIONS = {"doc", "docx", "md", "pdf", "txt"}


def load_local_env():
    """Load simple local development variables without adding another runtime dependency."""
    env_file = os.path.join(BASE_DIR, ".env")
    if not os.path.isfile(env_file):
        return
    with open(env_file, encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            name, value = line.split("=", 1)
            os.environ.setdefault(name.strip(), value.strip().strip('"').strip("'"))


load_local_env()
STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "in", "is",
    "it", "of", "on", "or", "that", "the", "this", "to", "was", "with", "will",
}
STOP_CONCEPTS = {
    "this", "that", "these", "those", "there", "here", "below", "above",
    "it", "they", "we", "you", "each", "every", "all", "some", "none",
    "one", "table", "page", "chapter", "section", "note", "figure", "habits",
    "organs", "vitamins", "minerals", "checklist", "foundations", "guide",
    "a complete", "the following", "key takeaway", "summary", "overview",
    "both", "both tags", "other", "another", "many", "such", "several", "most",
    "either", "neither", "various", "certain", "more", "less", "few", "any",
    "which", "whose", "what", "tags", "tag", "elements", "element", "attributes",
}

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("OMNISTUDY_SECRET_KEY", "change-this-before-production"),
    MAX_CONTENT_LENGTH=10 * 1024 * 1024,
    UPLOAD_FOLDER=UPLOAD_FOLDER,
)
scheduler = BackgroundScheduler(timezone="Asia/Kolkata")


def now_iso():
    return datetime.now(UTC).isoformat()


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_error):
    db = g.pop("db", None)
    if db is not None:
        db.close()


@app.after_request
def add_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


def ensure_column(db, table, column_definition):
    column_name = column_definition.split()[0]
    columns = {row["name"] for row in db.execute(f"PRAGMA table_info({table})")}
    if column_name not in columns:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {column_definition}")


def init_db():
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL DEFAULT '',
            password_hash TEXT,
            role TEXT NOT NULL CHECK(role IN ('student', 'staff', 'admin')),
            department TEXT,
            semester TEXT,
            last_login TEXT NOT NULL,
            is_verified INTEGER NOT NULL DEFAULT 1,
            graduation_status TEXT NOT NULL DEFAULT 'active'
        );
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            subject TEXT NOT NULL,
            department TEXT NOT NULL,
            semester TEXT NOT NULL,
            content TEXT NOT NULL,
            visibility TEXT NOT NULL CHECK(visibility IN ('public', 'private')),
            release_on_inactivity INTEGER NOT NULL DEFAULT 0,
            owner_id INTEGER NOT NULL,
            parent_id INTEGER,
            created_at TEXT NOT NULL,
            source_filename TEXT,
            stored_filename TEXT,
            processing_status TEXT NOT NULL DEFAULT 'ready',
            access_mode TEXT NOT NULL DEFAULT 'all',
            FOREIGN KEY(owner_id) REFERENCES users(id),
            FOREIGN KEY(parent_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS quiz_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS activity_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            document_id INTEGER,
            event_type TEXT NOT NULL,
            detail TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(document_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS study_summaries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            position INTEGER NOT NULL,
            sentence TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS quiz_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            position INTEGER NOT NULL,
            prompt TEXT NOT NULL,
            options_json TEXT NOT NULL,
            answer TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS verification_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            admin_id INTEGER NOT NULL,
            action TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(admin_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS document_access (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            permission TEXT NOT NULL CHECK(permission IN ('allow', 'block')),
            UNIQUE(document_id, user_id),
            FOREIGN KEY(document_id) REFERENCES documents(id),
            FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS system_settings (
            setting_key TEXT PRIMARY KEY,
            setting_value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS flashcard_progress (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            document_id INTEGER NOT NULL,
            card_key TEXT NOT NULL,
            box_level INTEGER NOT NULL DEFAULT 1,
            review_count INTEGER NOT NULL DEFAULT 0,
            last_reviewed TEXT,
            UNIQUE(user_id, document_id, card_key),
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(document_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS viva_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            document_id INTEGER NOT NULL,
            question TEXT NOT NULL,
            student_answer TEXT NOT NULL,
            score REAL NOT NULL,
            feedback TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(document_id) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS profile_update_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            current_name TEXT NOT NULL,
            new_name TEXT NOT NULL,
            current_reg_no TEXT NOT NULL DEFAULT '',
            new_reg_no TEXT NOT NULL DEFAULT '',
            current_department TEXT NOT NULL,
            new_department TEXT NOT NULL,
            current_semester TEXT NOT NULL,
            new_semester TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending' CHECK(status IN ('pending', 'approved', 'rejected')),
            admin_id INTEGER,
            created_at TEXT NOT NULL,
            reviewed_at TEXT,
            FOREIGN KEY(user_id) REFERENCES users(id),
            FOREIGN KEY(admin_id) REFERENCES users(id)
        );
        """
    )
    ensure_column(db, "users", "reg_no TEXT DEFAULT ''")
    ensure_column(db, "users", "password_hash TEXT")
    ensure_column(db, "users", "is_verified INTEGER NOT NULL DEFAULT 1")
    ensure_column(db, "users", "graduation_status TEXT NOT NULL DEFAULT 'active'")
    ensure_column(db, "documents", "source_filename TEXT")
    ensure_column(db, "documents", "stored_filename TEXT")
    ensure_column(db, "documents", "processing_status TEXT NOT NULL DEFAULT 'ready'")
    ensure_column(db, "documents", "access_mode TEXT NOT NULL DEFAULT 'all'")
    if not db.execute("SELECT 1 FROM users LIMIT 1").fetchone():
        demo_users = [
            ("Asha Professor", "staff@omnistudy.test", "staff123", "staff", "Computer Science", "All", ""),
            ("Ravi Student", "student@omnistudy.test", "student123", "student", "Computer Science", "6", "22BCA101"),
            ("System Admin", "admin@omnistudy.test", "admin123", "admin", "Administration", "All", ""),
        ]
        db.executemany(
            "INSERT INTO users (name, email, password, password_hash, role, department, semester, last_login, reg_no) VALUES (?, ?, '', ?, ?, ?, ?, ?, ?)",
            [(name, email, generate_password_hash(password), role, department, semester, now_iso(), reg_no) for name, email, password, role, department, semester, reg_no in demo_users],
        )
    db.execute("UPDATE users SET reg_no = '22BCA101' WHERE email = 'student@omnistudy.test' AND (reg_no IS NULL OR reg_no = '')")
    for user in db.execute("SELECT id, password, password_hash FROM users WHERE password_hash IS NULL OR password_hash = ''"):
        db.execute("UPDATE users SET password_hash = ?, password = '' WHERE id = ?", (generate_password_hash(user["password"]), user["id"]))
    db.executemany(
        "INSERT OR IGNORE INTO system_settings (setting_key, setting_value) VALUES (?, ?)",
        [("inactivity_days", "90"), ("last_watchdog_run", "Not run yet")],
    )
    db.commit()


@app.before_request
def load_user():
    if "_csrf_token" not in session:
        session["_csrf_token"] = secrets.token_urlsafe(32)
    user_id = session.get("user_id")
    g.user = get_db().execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone() if user_id else None


@app.context_processor
def csrf_context():
    token = session.setdefault("_csrf_token", secrets.token_urlsafe(32))
    return {"csrf_token": token}


@app.before_request
def protect_csrf():
    if app.config.get("WTF_CSRF_ENABLED") is False:
        return
    if request.method == "POST":
        expected_token = session.get("_csrf_token", "")
        submitted_token = (
            request.form.get("_csrf_token", "")
            or request.headers.get("X-CSRFToken", "")
            or request.headers.get("X-CSRF-Token", "")
        )
        if not submitted_token and request.is_json:
            data = request.get_json(silent=True)
            if isinstance(data, dict):
                submitted_token = data.get("_csrf_token", "")
        if not expected_token or not submitted_token or not hmac.compare_digest(expected_token, submitted_token):
            abort(400)


def login_required(roles=None):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if g.user is None:
                return redirect(url_for("login"))
            if roles and g.user["role"] not in roles:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def log_activity(event_type, document_id=None, detail=None):
    get_db().execute(
        "INSERT INTO activity_logs (user_id, document_id, event_type, detail, created_at) VALUES (?, ?, ?, ?, ?)",
        (g.user["id"], document_id, event_type, detail, now_iso()),
    )
    get_db().commit()


def visible_document(document):
    if not document:
        return False
    if not g.user:
        return document["visibility"] == "public"
    if g.user["role"] == "admin" or document["owner_id"] == g.user["id"]:
        return True
    access_mode = document["access_mode"] or "all"
    if access_mode == "all":
        return document["visibility"] == "public"
    if access_mode == "allow_list":
        return g.user["role"] == "student" and get_db().execute(
            "SELECT 1 FROM document_access WHERE document_id = ? AND user_id = ? AND permission = 'allow'",
            (document["id"], g.user["id"]),
        ).fetchone() is not None
    if access_mode == "block_list":
        return document["visibility"] == "public" and get_db().execute(
            "SELECT 1 FROM document_access WHERE document_id = ? AND user_id = ? AND permission = 'block'",
            (document["id"], g.user["id"]),
        ).fetchone() is None
    return False


def available_students():
    return get_db().execute(
        "SELECT id, name, email, department, semester FROM users WHERE role = 'student' AND is_verified = 1 ORDER BY name"
    ).fetchall()


def selected_student_ids(document_id, permission):
    return [row["user_id"] for row in get_db().execute(
        "SELECT user_id FROM document_access WHERE document_id = ? AND permission = ?", (document_id, permission)
    )]


def validate_document_access(access_mode, requested_student_ids):
    if access_mode not in {"all", "allow_list", "block_list"}:
        raise ValueError("Select a valid student access option.")
    db = get_db()
    student_ids = []
    for student_id in requested_student_ids:
        try:
            student_ids.append(int(student_id))
        except ValueError:
            continue
    if access_mode in {"allow_list", "block_list"} and not student_ids:
        raise ValueError("Select at least one student for this access rule.")
    valid_ids = {
        row["id"] for row in db.execute(
            "SELECT id FROM users WHERE id IN ({}) AND role = 'student' AND is_verified = 1".format(",".join("?" for _ in student_ids) or "NULL"),
            student_ids,
        )
    }
    if set(student_ids) != valid_ids:
        raise ValueError("One or more selected accounts are not verified student accounts.")
    return list(valid_ids)


def save_document_access(document_id, access_mode, requested_student_ids):
    student_ids = validate_document_access(access_mode, requested_student_ids)
    permission = "allow" if access_mode == "allow_list" else "block"
    db = get_db()
    db.execute("UPDATE documents SET access_mode = ? WHERE id = ?", (access_mode, document_id))
    db.execute("DELETE FROM document_access WHERE document_id = ?", (document_id,))
    if access_mode != "all":
        db.executemany(
            "INSERT INTO document_access (document_id, user_id, permission) VALUES (?, ?, ?)",
            [(document_id, student_id, permission) for student_id in student_ids],
        )
    db.commit()


def accessible_documents(search_term=""):
    sql = "SELECT d.*, u.name AS owner_name FROM documents d JOIN users u ON u.id = d.owner_id"
    parameters = []
    if search_term:
        sql += " WHERE d.title LIKE ? OR d.subject LIKE ? OR d.content LIKE ?"
        parameters.extend([f"%{search_term}%"] * 3)
    sql += " ORDER BY d.created_at DESC"
    return [document for document in get_db().execute(sql, parameters).fetchall() if visible_document(document)]


def get_document(document_id):
    return get_db().execute(
        "SELECT d.*, u.name AS owner_name FROM documents d JOIN users u ON u.id = d.owner_id WHERE d.id = ?", (document_id,)
    ).fetchone()


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_upload(upload):
    if not upload or not upload.filename:
        return "", None, None
    if not allowed_file(upload.filename):
        raise ValueError("Only PDF, Word (.docx), TXT, and Markdown files are supported.")
    original_ext = upload.filename.rsplit(".", 1)[1].lower() if "." in upload.filename else ""
    source_filename = secure_filename(upload.filename)
    if not source_filename or "." not in source_filename:
        source_filename = f"upload_{uuid.uuid4().hex[:8]}.{original_ext}"
    stored_filename = f"{uuid.uuid4().hex}_{source_filename}"
    stored_path = os.path.join(app.config["UPLOAD_FOLDER"], stored_filename)
    upload.save(stored_path)
    extension = original_ext
    try:
        if extension == "pdf":
            if PdfReader is None:
                raise ValueError("PDF support is unavailable. Install the project dependencies and try again.")
            content = "\n".join(page.extract_text() or "" for page in PdfReader(stored_path).pages)
        elif extension in {"docx", "doc"}:
            if docx is not None:
                doc = docx.Document(stored_path)
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            paragraphs.append(" | ".join(row_cells))
                content = "\n\n".join(paragraphs)
            else:
                # Fallback native DOCX extraction via zipfile
                with zipfile.ZipFile(stored_path) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    texts = [node.text for node in tree.iter() if node.tag.endswith("}t") and node.text]
                    content = " ".join(texts)
        else:
            with open(stored_path, "r", encoding="utf-8", errors="ignore") as text_file:
                content = text_file.read()
    except Exception as error:
        if os.path.exists(stored_path):
            os.remove(stored_path)
        raise ValueError(f"Could not read the uploaded file: {error}") from error
    if not content.strip():
        if os.path.exists(stored_path):
            os.remove(stored_path)
        raise ValueError("The uploaded file contains no readable text.")
    return content.strip(), source_filename, stored_filename


def clean_study_text(raw_text):
    """Strip page headers, page numbers, table of contents noise, bullets, and unwrap broken lines."""
    lines = raw_text.splitlines()
    cleaned = []
    for line in lines:
        l = line.strip()
        if not l:
            continue
        if re.search(r"^page\s+\d+$", l, re.IGNORECASE):
            continue
        if re.search(r"table of contents", l, re.IGNORECASE):
            continue
        if re.search(r"^a complete guide to", l, re.IGNORECASE):
            continue
        if re.search(r"^habits\s+organs\s+vitamins", l, re.IGNORECASE):
            continue
        if re.search(r"^\d+\.\s+(foundations|organ-by-organ|vitamins|minerals|hydration|warning|checklist)", l, re.IGNORECASE) and len(l) < 55:
            continue
        l = re.sub(r"^[■❤️❤★✓•\-\–\—\*\s]+", "", l).strip()
        l = re.sub(r"[■❤️❤★✓•]", "", l).strip()
        if len(l) > 2:
            cleaned.append(l)

    # Unwrap broken lines into coherent full sentences
    merged_lines = []
    for line in cleaned:
        if merged_lines and not merged_lines[-1].endswith((".", "!", "?", ":")) and not re.match(r"^[A-Z0-9\s&/\-\(\)]{3,40}$", line) and not line.lower().startswith(("role", "habit", "watch", "tip", "note")):
            merged_lines[-1] += " " + line
        else:
            merged_lines.append(line)

    return "\n".join(merged_lines)


def extract_clean_sentences(content):
    """Cleanly split study text into high-quality, distinct academic sentences without merging table noise."""
    lines = content.splitlines()
    candidate_sentences = []
    
    for raw_line in lines:
        line = raw_line.strip()
        if not line or len(line) < 20:
            continue
        if re.search(r"[─\-_=]{3,}|[►◄→←|]", line) and line.count("|") > 1:
            continue
        if re.search(r"^(diagram|table of contents|exam tip|draw this|page\s+\d+|topic\s+\d+)", line, re.IGNORECASE):
            continue
        if line.startswith("•") or line.startswith("-") or line.startswith("*"):
            line = line.lstrip("•-* ").strip()
            
        chunks = re.split(r"(?<=[.!?])\s+", line)
        for chunk in chunks:
            c = chunk.strip()
            c = re.sub(r"^\d+[\)\.]\s*", "", c)
            c = re.sub(r'^(one-line\s+exam\s+answer|simple\s+definition\s+line\s+for\s+exam):\s*', "", c, flags=re.IGNORECASE)
            c = c.strip('"\': ')
            
            if 35 <= len(c) <= 220:
                if re.search(r"\b(is|are|provides|manages|enables|supports|acts|helps|stores|organizes|connects|defines|describes|contains|creates|solves|ensures|regulates|maintains|filters|processes)\b", c, re.IGNORECASE):
                    candidate_sentences.append(c)
                    
    return candidate_sentences


def summarize(content, count=5):
    candidates = extract_clean_sentences(content)
    if not candidates:
        cleaned = clean_study_text(content)
        full_text = " ".join(cleaned.splitlines())
        sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", full_text) if len(s.strip()) > 30 and not s.lower().startswith(("page", "table", "a complete"))]
        if not sentences:
            return [content.strip()[:200]] if content.strip() else []
        frequency = Counter(word for word in re.findall(r"[a-zA-Z]{3,}", cleaned.lower()) if word not in STOP_WORDS)
        ranked = sorted(sentences, key=lambda sentence: sum(frequency[word] for word in re.findall(r"[a-zA-Z]{3,}", sentence.lower())), reverse=True)
        return ranked[:count]
        
    words = [w.lower() for w in re.findall(r"[a-zA-Z]{3,}", content) if w.lower() not in STOP_WORDS]
    freq = Counter(words)
    seen_prefixes = set()
    scored = []
    
    for s in candidates:
        prefix = s[:30].lower()
        if prefix in seen_prefixes:
            continue
        seen_prefixes.add(prefix)
        score = sum(freq[w.lower()] for w in re.findall(r"[a-zA-Z]{3,}", s))
        if re.search(r"\b(is\s+(?:the|a|an)?|refers\s+to|defines|provides|enables)\b", s, re.IGNORECASE):
            score *= 1.3
        scored.append((score, s))
        
    scored.sort(key=lambda x: x[0], reverse=True)
    return [s for _, s in scored[:count]]


DISTRACTOR_BANKS = {
    "cs": {
        "mechanisms": [
            "Computes cryptographic hash digests of keys to facilitate near constant-time associative lookup.",
            "Recursively partitions array elements around selected pivots to enforce comparative order.",
            "Traverses hierarchically structured nodes via breadth-first search using double-ended queues.",
            "Constructs balanced binary search trees to guarantee logarithmic lookup, insertion, and deletion times.",
            "Applies memoization across overlapping subproblems to optimize exponential recursive time complexity.",
            "Maintains transactional atomicity and strict serializability using two-phase locking protocols.",
            "Coordinates asynchronous data transmission and message buffering across decoupled network sockets.",
            "Facilitates automatic garbage collection and memory compaction through generational mark-and-sweep.",
            "Executes dynamic branch prediction and out-of-order instruction scheduling to maximize throughput.",
            "Translates intermediate bytecode representations into native machine instructions via JIT compilation.",
            "Enforces role-based access control policies and validates cryptographic certificates during handshakes.",
            "Synchronizes thread access to critical sections using reentrant mutex locks and atomic operations.",
            "Distributes incoming web traffic across redundant application clusters using weighted round-robin.",
            "Compacts inverted search indexes to facilitate high-performance full-text semantic document retrieval.",
            "Enforces relational referential integrity by validating foreign key references across linked tables.",
            "Organizes records into tabular structures of rows and columns to simplify declarative querying.",
            "Defines document structure and content semantics using standardized paired markup elements.",
            "Applies cascading style rules to decouple presentation semantics from underlying markup.",
            "Manages asynchronous events and server responses using single-threaded event loops.",
            "Maintains data consistency and atomicity across concurrent multi-user database transactions.",
        ],
        "concepts": [
            "Binary Search Tree", "Dynamic Hash Index", "Red-Black Tree", "Priority Heap Queue",
            "Direct Memory Access Controller", "Interrupt Service Routine", "Virtual Memory Paging Unit",
            "Transaction Coordinator", "Distributed Consensus Engine", "Cryptographic Keystore",
            "Asynchronous Event Loop", "JIT Compiler Optimizer", "Load Balancer Proxy",
            "Inverted Index Store", "Two-Phase Commit Manager", "Deadlock Detection Monitor",
            "Divide and Conquer Strategy", "Asymptotic Complexity Estimator", "Dynamic Memory Allocator",
            "Database Management System", "Relational Data Model", "Primary Key Index",
            "Foreign Key Constraint", "Entity-Relationship Model", "Document Object Model",
            "Semantic Markup Element", "Cascading Style Sheet", "Client-Side Router",
        ],
        "definitions": [
            "An organized collection of structured data stored electronically in a computational system.",
            "A standardized markup language used to structure web documents and delineate content meaning.",
            "A unique attribute or set of attributes that unambiguously identifies a tuple within a relation.",
            "A declarative query language designed for managing and querying structured relational databases.",
            "A programming paradigm based on the concept of objects containing data fields and procedures.",
            "A software subsystem that manages data storage, retrieval, and concurrency control.",
            "A structural constraint ensuring valid relationships and referential integrity across tables.",
        ],
        "systems": [
            "Virtual Memory Subsystem", "Process Scheduling Subsystem", "Network Protocol Stack",
            "Database Storage Engine", "File Allocation System", "Cryptographic Security Module",
            "Distributed Consensus Cluster", "Thread Execution Manager",
        ],
    },
    "biomedical": {
        "mechanisms": [
            "Filters metabolic waste products and excess urea from plasma while maintaining systemic electrolyte balance.",
            "Facilitates pulmonary alveolar gas exchange by absorbing molecular oxygen and expelling carbon dioxide.",
            "Pumps oxygenated erythrocytes through high-pressure arterial networks to sustain peripheral tissue perfusion.",
            "Secretes pancreatic endocrine hormones including insulin and glucagon to govern systemic glucose homeostasis.",
            "Synthesizes essential bile acids, metabolizes dietary macronutrients, and neutralizes circulating xenobiotics.",
            "Coordinates afferent sensory input and efferent motor responses via rapid electrochemical neurotransmission.",
            "Maintains structural skeletal rigidity, protects visceral organs, and reservoirs calcium and phosphate ions.",
            "Generates contractile mechanical force and metabolic thermogenesis via actin-myosin cross-bridge cycling.",
            "Provides stratified epithelial barrier defenses and regulates transdermal evaporative thermoregulation.",
            "Mounts adaptive humoral antibody responses and cell-mediated cytotoxic defense against foreign antigens.",
            "Absorbs dietary lipids and hydrolyzes macronutrients along the microvillar epithelial border.",
            "Regulates arterial baroreceptor reflexes and autonomic sympathetic vascular vasomotor tone.",
        ],
        "concepts": [
            "Cardiovascular & Hemodynamic System", "Pulmonary & Respiratory System", "Central & Peripheral Nervous System",
            "Hepatic & Biliary System", "Renal & Urinary System", "Endocrine & Pancreatic Axis",
            "Integumentary & Barrier System", "Musculoskeletal & Articular System", "Gastrointestinal & Enteric System",
            "Lymphatic & Adaptive Immune System",
        ],
        "systems": [
            "Cardiovascular & Circulatory System", "Pulmonary & Respiratory System", "Central & Peripheral Nervous System",
            "Hepatic & Biliary System", "Renal & Urinary System", "Endocrine & Pancreatic Axis",
            "Integumentary & Barrier System", "Musculoskeletal & Articular System", "Gastrointestinal & Enteric System",
            "Lymphatic & Immune Defense System",
        ],
        "concise_organs": [
            "Heart & Circulation", "Lungs & Airway", "Brain & Nerves", "Liver & Gallbladder",
            "Kidneys & Bladder", "Pancreas & Spleen", "Stomach & Intestines", "Skin & Integument",
            "Bones & Joints", "Muscles & Tendons", "Eyes & Retinal Axis",
        ],
        "nutrients": [
            "Vitamin A (Retinol)", "Vitamin B1 (Thiamine)", "Vitamin B2 (Riboflavin)", "Vitamin B3 (Niacin)",
            "Vitamin B6 (Pyridoxine)", "Vitamin B9 (Folate)", "Vitamin B12 (Cobalamin)", "Vitamin C (Ascorbic Acid)",
            "Vitamin D3 (Cholecalciferol)", "Vitamin E (Tocopherol)", "Vitamin K2 (Menaquinone)", "Ionic Calcium (Ca2+)",
            "Heme Iron (Fe2+)", "Magnesium (Mg2+)", "Zinc (Zn2+)", "Potassium (K+)",
        ],
        "nutrient_roles": [
            "Catalyzes mitochondrial electron transport chain complexes to generate cellular ATP.",
            "Maintains bone mineral density by regulating parathyroid-dependent intestinal calcium absorption.",
            "Serves as an essential cofactor in hepatic coagulation factor synthesis (II, VII, IX, X).",
            "Scavenges reactive oxygen species to protect polyunsaturated phospholipid membranes from peroxidation.",
            "Facilitates reversible oxygen binding and systemic tissue delivery within erythrocyte hemoglobin complexes.",
            "Serves as an indispensable methyl donor cofactor in genomic DNA methylation and homocysteine metabolism.",
            "Maintains intracellular osmotic potential and drives neuronal resting membrane potential repolarization.",
            "Supports thymocyte differentiation, T-lymphocyte maturation, and mucosal barrier integrity.",
        ],
        "concise_habits": [
            "Nutritious Whole Foods", "Daily Physical Activity", "Mindful Stress Reduction",
            "Regular Medical Check-Ups", "Restorative Sleep Hygiene", "Optimal Hydration Intake",
            "Toxin & Tobacco Avoidance", "Cognitive & Mental Stimulation", "Posture & Mobility Work",
        ],
        "habits": [
            "Consuming Nutrient-Dense Whole Foods & Lean Proteins",
            "Engaging in Daily Structured Aerobic & Resistance Activity",
            "Practicing Proactive Stress Modulation & Mindfulness Protocols",
            "Adhering to Regular Preventative Health Screenings & Check-Ups",
            "Maintaining Consistent Circadian Sleep Hygiene & Architecture",
            "Sustaining Optimal Daily Hydration & Electrolyte Homeostasis",
            "Limiting Exposure to Environmental Toxins & Ultra-Processed Food",
            "Engaging in Lifelong Cognitive Enrichment & Neurological Training",
        ],
        "definitions": [
            "A physiological mechanism that maintains internal systemic equilibrium despite external fluctuations.",
            "A microscopic functional unit responsible for specialized metabolic and cellular processing.",
            "An organic compound required in micronutrient quantities to sustain vital enzymatic reactions.",
        ],
    },
    "general": {
        "mechanisms": [
            "Establishes empirical measurement criteria to evaluate structural hypothesis validity across varied test conditions.",
            "Coordinates resource allocation models to maximize operational efficiency under constrained operational bounds.",
            "Maintains regulatory compliance frameworks through continuous audit logging and anomaly detection protocols.",
            "Optimizes functional throughput by eliminating procedural bottlenecks and standardizing core workflows.",
            "Synthesizes multidimensional feedback loops to support adaptive organizational decision-making.",
            "Enforces quantitative benchmark thresholds to ensure consistent quality assurance and defect minimization.",
        ],
        "concepts": [
            "Empirical Methodology Framework", "Operational Constraint Model", "Systemic Feedback Architecture",
            "Dynamic Optimization Protocol", "Quantitative Benchmark Metric", "Adaptive Governance Mechanism",
            "Resource Allocation Engine", "Continuous Verification Model",
        ],
        "definitions": [
            "A structured analytical methodology designed for reproducible observation and verification.",
            "An operational set of constraints ensuring standardized performance and measurable output.",
            "A comprehensive framework governing quality assurance and consistent policy execution.",
        ],
    },
}


def detect_domain(text):
    t_low = text.lower()
    cs_keywords = ["algorithm", "search", "tree", "memory", "database", "network", "cache", "server", "cpu", "gpu", "compiler", "queue", "stack", "hash", "code", "thread", "process", "byte", "file", "software", "api", "query", "html", "tag", "css", "rdbms", "sql"]
    bio_keywords = ["body", "organ", "blood", "heart", "lung", "liver", "kidney", "brain", "vitamin", "nutrient", "cell", "respiratory", "cardiovascular", "stomach", "skin", "muscle", "bone", "diet", "sleep", "health", "symptom", "disease"]
    cs_count = sum(t_low.count(kw) for kw in cs_keywords)
    bio_count = sum(t_low.count(kw) for kw in bio_keywords)
    if cs_count > bio_count and cs_count > 0:
        return "cs"
    if bio_count >= cs_count and bio_count > 0:
        return "biomedical"
    return "general"


def build_smart_options(correct_answer, candidate_distractors, fallback_pool):
    """Build exactly 4 distinct, grammatically matched, high-plausibility options."""
    correct_clean = correct_answer.strip()
    target_len = len(correct_clean)
    seen_lower = {correct_clean.lower()}
    chosen = []

    def clean_and_validate(item):
        if not item:
            return None
        c = str(item).strip()
        if not c or c.lower() in seen_lower:
            return None
        # Disallow near-identical substrings (e.g. "Both" vs "Both Tags")
        if len(c) <= 7 and c.lower() in correct_clean.lower():
            return None
        if len(correct_clean) <= 7 and correct_clean.lower() in c.lower():
            return None
        # Length compatibility: avoid severe giveaways
        if target_len > 70 and len(c) < 22:
            return None
        if target_len < 25 and len(c) > 75:
            return None
        # Grammatical period alignment
        if correct_clean.endswith(".") and not c.endswith("."):
            c = c + "."
        elif not correct_clean.endswith(".") and c.endswith("."):
            c = c.rstrip(".")
        # Capitalization alignment
        if correct_clean and correct_clean[0].isupper() and c and c[0].islower():
            c = c[0].upper() + c[1:]
        if c.lower() in seen_lower:
            return None
        return c

    # 1. Candidate distractors sorted by length similarity
    scored_candidates = []
    for c in candidate_distractors:
        c_clean = clean_and_validate(c)
        if c_clean:
            diff = abs(len(c_clean) - target_len)
            scored_candidates.append((diff, c_clean))
    scored_candidates.sort(key=lambda x: x[0])
    for _, c_clean in scored_candidates:
        if len(chosen) >= 3:
            break
        if c_clean.lower() not in seen_lower:
            seen_lower.add(c_clean.lower())
            chosen.append(c_clean)

    # 2. Fallback pool sorted by length similarity
    scored_fallbacks = []
    for fb in fallback_pool:
        fb_clean = clean_and_validate(fb)
        if fb_clean:
            diff = abs(len(fb_clean) - target_len)
            scored_fallbacks.append((diff, fb_clean))
    scored_fallbacks.sort(key=lambda x: x[0])
    for _, fb_clean in scored_fallbacks:
        if len(chosen) >= 3:
            break
        if fb_clean.lower() not in seen_lower:
            seen_lower.add(fb_clean.lower())
            chosen.append(fb_clean)

    # 3. Final safety guarantee if pool exhausted
    counter = 1
    while len(chosen) < 3:
        if correct_clean.endswith("."):
            alt = f"Standard operational property {counter}."
        else:
            alt = f"Standard Operational Component {counter}"
        if alt.lower() not in seen_lower:
            seen_lower.add(alt.lower())
            chosen.append(alt)
        counter += 1

    options = [correct_clean] + chosen[:3]
    random.shuffle(options)
    return options


def make_quiz(raw_content, max_questions=30):
    cleaned_text = clean_study_text(raw_content)
    lines = cleaned_text.splitlines()
    domain = detect_domain(cleaned_text)
    banks = DISTRACTOR_BANKS.get(domain, DISTRACTOR_BANKS["general"])
    bio_banks = DISTRACTOR_BANKS["biomedical"]
    cs_banks = DISTRACTOR_BANKS["cs"]
    
    questions = []
    seen_prompts = set()

    # Pre-clean lines to avoid merging bullet headers across lines
    clean_lines = []
    for line in lines:
        l_str = line.strip()
        if not l_str:
            continue
        l_sub = re.sub(r"^[\s•\-\*–—\d\.\)\(]+", "", l_str).strip()
        if l_sub:
            if not l_sub[-1] in ".!?:;":
                clean_lines.append(l_sub + ".")
            else:
                clean_lines.append(l_sub)
    full_text = " ".join(clean_lines)

    # 1. Structure Extraction: "Topic / Header" followed by "Role / Function: ..."
    topic_roles = []
    topic_warnings = []
    current_topic = None
    for i, line in enumerate(lines):
        header_m = re.match(r"^([A-Z0-9\s&/\-\(\)]{3,40})$", line)
        if header_m and not line.lower().startswith(("role", "habit", "watch", "page", "tip", "note", "source", "below", "daily habits")):
            t = header_m.group(1).strip().title()
            if len(t.split()) <= 6 and not any(t.lower().startswith(sc) for sc in STOP_CONCEPTS):
                current_topic = t
            continue

        role_m = re.match(r"^Role(?:\s+in\s+the\s+body)?:\s*(.+)$", line, re.IGNORECASE)
        if role_m and current_topic:
            role_text = role_m.group(1).strip().rstrip(".")
            if len(role_text) > 8:
                topic_roles.append({"topic": current_topic, "role": role_text})

        watch_m = re.match(r"^Watch\s+out\s+for:\s*(.+)$", line, re.IGNORECASE)
        if watch_m and current_topic:
            watch_text = watch_m.group(1).strip().rstrip(".")
            if len(watch_text) > 8:
                topic_warnings.append({"topic": current_topic, "warning": watch_text})

    # 2. Key-Value Item & Definition Extraction (supports bullets and standard formats)
    item_definitions = []
    for line in lines:
        line_clean = re.sub(r"^[\s•\-\*–—\d\.\)\(]+", "", line).strip()
        kv_m = re.match(r"^([A-Z][a-zA-Z0-9\s\(\)\-\.,/]{2,35}?)\s+(?:—|–|:)\s+([A-Z0-9][a-zA-Z0-9\s\(\)\-\.,;:'\"/]+)$", line_clean)
        if kv_m:
            term = kv_m.group(1).strip().title()
            desc = kv_m.group(2).strip().rstrip(".")
            if (term.lower() not in STOP_CONCEPTS
                and not any(term.lower() == sc or term.lower().startswith(sc + " ") for sc in STOP_CONCEPTS)
                and len(term.split()) <= 5
                and len(desc) > 12):
                item_definitions.append({"term": term, "desc": desc, "raw": line})

    # 3. Nutrient / Biomedical Table Component Extraction (if biomedical)
    structured_nutrients = []
    if domain == "biomedical":
        nutrient_names = [
            "Vitamin A", "Vitamin B1 (Thiamine)", "Vitamin B2 (Riboflavin)", "Vitamin B3 (Niacin)",
            "Vitamin B6", "Vitamin B9 (Folate)", "Vitamin B12", "Vitamin C", "Vitamin D",
            "Vitamin E", "Vitamin K", "Calcium", "Iron", "Magnesium", "Zinc", "Potassium"
        ]
        for idx, l in enumerate(lines):
            for n in nutrient_names:
                if l.lower().startswith(n.lower()):
                    remainder = l[len(n):].strip().lstrip("—–:- ").strip().rstrip(".")
                    if len(remainder) > 10:
                        structured_nutrients.append({"name": n, "role": remainder})
                    elif idx + 1 < len(lines):
                        r_text = lines[idx+1].strip().rstrip(".")
                        if len(r_text) > 10 and not any(r_text.lower().startswith(sc) for sc in STOP_CONCEPTS) and not any(r_text.lower().startswith(x.lower()) for x in nutrient_names):
                            structured_nutrients.append({"name": n, "role": r_text})
                    break

    # 4. Comprehensive Sentence-level Concept Extraction
    raw_sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", full_text) if len(s.strip()) > 15]
    academic_verbs = (
        r"is\s+(?:the|a|an)?|are|provides?|manages?|enables?|determines?|organizes?|prevents?|"
        r"occurs\s+when|refers\s+to|regulates?|supports?|facilitates?|clears?|protects?|solves?|"
        r"works?\s+on|reduces?|helps?|maps?|extends?|translates?|generates?|controls?|allows?|"
        r"stores?|executes?|processes?|performs?|coordinates?|implements?|maintains?|synthesizes?|"
        r"filters?|converts?|transforms?|measures?|calculates?|estimates?|optimizes?|structures?|partitions?|"
        r"defines?|represents?|contains?|includes?|specifies?|ensures?|establishes?|guarantees?"
    )
    sentence_pattern = rf"^(?:The\s+|A\s+|An\s+)?([A-Z][a-zA-Z0-9_\-\s]{{2,35}}?)\s+({academic_verbs})\s+(.+)$"

    concept_sentences = []
    for s in raw_sentences:
        m = re.match(sentence_pattern, s, re.IGNORECASE)
        if m:
            term = m.group(1).strip().title()
            if term.lower().endswith((" can", " may", " must", " should", " will", " could")):
                term = term.rsplit(" ", 1)[0].strip()
            verb = m.group(2).strip()
            pred = m.group(3).strip().rstrip(".")
            if (term.lower() not in STOP_CONCEPTS and 
                not any(term.lower() == sc or term.lower().startswith(sc + " ") for sc in STOP_CONCEPTS) and 
                " is when " not in term.lower() and 
                len(term.split()) <= 4 and 
                len(pred) > 8):
                concept_sentences.append({"term": term, "verb": verb, "predicate": pred, "sentence": s})

    all_roles = [tr["role"] for tr in topic_roles]
    all_nutrient_roles = [sn["role"] for sn in structured_nutrients]
    other_concept_terms = [cs["term"] for cs in concept_sentences]
    other_item_terms = [it["term"] for it in item_definitions]

    # Question Type 1: System / Component Primary Function (Easy / Medium)
    for tr in topic_roles:
        if len(questions) >= max_questions:
            break
        prompt = f"What is the primary function or operational role of \"{tr['topic']}\"?"
        if prompt not in seen_prompts:
            seen_prompts.add(prompt)
            correct_ans = tr["role"][0].upper() + tr["role"][1:] + "."
            other_roles = [r[0].upper() + r[1:] + "." for r in all_roles if r != tr["role"]]
            opts = build_smart_options(correct_ans, other_roles, banks.get("mechanisms", bio_banks["mechanisms"]))
            questions.append({
                "prompt": prompt,
                "options": opts,
                "answer": correct_ans,
                "explanation": f"According to the study material, the primary role of {tr['topic']} is: {tr['role']}.",
                "skill": "System Function & Analysis",
                "difficulty": "Easy" if len(questions) % 2 == 0 else "Medium"
            })

    # Question Type 2: Clinical Warning Signs & Diagnostic Indicators (Medium) - Biomedical only
    if domain == "biomedical":
        for tw in topic_warnings:
            if len(questions) >= max_questions:
                break
            prompt = f"Warning signs including \"{tw['warning']}\" are key clinical indicators of potential dysfunction in which organ or system?"
            if prompt not in seen_prompts:
                seen_prompts.add(prompt)
                correct_ans = tw["topic"]
                fallback_organs = bio_banks["systems"] if ("&" in correct_ans or "System" in correct_ans) else bio_banks["concise_organs"]
                other_topics = [t for t in [tr["topic"] for tr in topic_roles] if t != correct_ans]
                opts = build_smart_options(correct_ans, other_topics, fallback_organs)
                questions.append({
                    "prompt": prompt,
                    "options": opts,
                    "answer": correct_ans,
                    "explanation": f"The study material states to watch out for '{tw['warning']}' regarding the health of the {tw['topic']}.",
                    "skill": "Diagnostic Indicators & Pathology",
                    "difficulty": "Medium"
                })

    # Question Type 3: Vitamin & Mineral Specific Roles (Easy / Medium)
    for sn in structured_nutrients:
        if len(questions) >= max_questions:
            break
        prompt = f"What is the primary physiological function or health role of \"{sn['name']}\"?"
        if prompt not in seen_prompts:
            seen_prompts.add(prompt)
            correct_ans = sn["role"][0].upper() + sn["role"][1:] + "."
            other_roles = [r[0].upper() + r[1:] + "." for r in all_nutrient_roles if r != sn["role"]]
            opts = build_smart_options(correct_ans, other_roles, bio_banks["nutrient_roles"])
            questions.append({
                "prompt": prompt,
                "options": opts,
                "answer": correct_ans,
                "explanation": f"According to the nutrient study guide, {sn['name']} is responsible for: {sn['role']}.",
                "skill": "Nutrient Biochemistry & Dietetics",
                "difficulty": "Easy" if len(questions) % 2 == 0 else "Medium"
            })

    # Question Type 4: Key-Value Item Identification (Easy / Medium)
    for idef in item_definitions:
        if len(questions) >= max_questions:
            break
        prompt = f"Which concept or component is characterized by: \"{idef['desc']}\"?"
        if prompt not in seen_prompts:
            seen_prompts.add(prompt)
            correct_ans = idef["term"]
            other_items = [t["term"] for t in item_definitions if t["term"] != correct_ans and t["term"].lower() not in correct_ans.lower()]
            fallback_items = banks.get("concepts", cs_banks["concepts"])
            opts = build_smart_options(correct_ans, other_items, fallback_items)
            questions.append({
                "prompt": prompt,
                "options": opts,
                "answer": correct_ans,
                "explanation": f"The study guide associates {idef['term']} with: {idef['desc']}.",
                "skill": "Principle Identification",
                "difficulty": "Easy" if len(questions) % 2 == 0 else "Medium"
            })

    # Question Type 4B: Key-Value Description Matching (Medium)
    for idef in item_definitions:
        if len(questions) >= max_questions:
            break
        prompt = f"How is \"{idef['term']}\" defined or described in the study material?"
        if prompt not in seen_prompts:
            seen_prompts.add(prompt)
            correct_ans = idef["desc"][0].upper() + idef["desc"][1:] + "."
            other_descs = [t["desc"][0].upper() + t["desc"][1:] + "." for t in item_definitions if t["term"] != idef["term"]]
            fallback_descs = banks.get("definitions", banks.get("mechanisms", cs_banks["mechanisms"]))
            opts = build_smart_options(correct_ans, other_descs, fallback_descs)
            questions.append({
                "prompt": prompt,
                "options": opts,
                "answer": correct_ans,
                "explanation": f"The study material defines {idef['term']}: {idef['desc']}.",
                "skill": "Conceptual Knowledge",
                "difficulty": "Medium"
            })

    # Question Type 5: Sentence Action & Mechanism (Easy / Medium / Hard)
    for cs in concept_sentences:
        if len(questions) >= max_questions:
            break
        is_definitional = cs['verb'].lower().startswith(("is", "are", "refers", "defines", "represents"))
        if is_definitional:
            prompt = f"How is \"{cs['term']}\" defined or characterized in the study material?"
            correct_ans = cs['predicate'][0].upper() + cs['predicate'][1:] + "."
            other_preds = [other['predicate'][0].upper() + other['predicate'][1:] + "." for other in concept_sentences if other['term'] != cs['term']]
            fallback_preds = banks.get("definitions", banks.get("mechanisms", cs_banks["mechanisms"]))
            difficulty = "Easy" if len(questions) % 2 == 0 else "Medium"
            skill = "Foundational Definition"
        else:
            prompt = f"What is the primary operational action, mechanism, or principle associated with \"{cs['term']}\"?"
            correct_ans = f"{cs['verb'].capitalize()} {cs['predicate']}."
            other_preds = [f"{other['verb'].capitalize()} {other['predicate']}." for other in concept_sentences if other['term'] != cs['term']]
            fallback_preds = banks.get("mechanisms", cs_banks["mechanisms"])
            difficulty = "Hard"
            skill = "Principle Application"

        if prompt not in seen_prompts:
            seen_prompts.add(prompt)
            opts = build_smart_options(correct_ans, other_preds, fallback_preds)
            questions.append({
                "prompt": prompt,
                "options": opts,
                "answer": correct_ans,
                "explanation": f"The study text confirms: {cs['sentence']}",
                "skill": skill,
                "difficulty": difficulty
            })

    # Question Type 6: Conceptual Identification
    if len(questions) < max_questions and concept_sentences:
        for cs in concept_sentences:
            if len(questions) >= max_questions:
                break
            if cs['verb'].lower() in {"are", "include"}:
                prompt = f"Which concept or component is characterized by: \"{cs['verb'].lower()} {cs['predicate']}\"?"
            else:
                prompt = f"Which concept or component {cs['verb'].lower()} {cs['predicate']}?"
            if prompt not in seen_prompts:
                seen_prompts.add(prompt)
                correct_ans = cs["term"]
                other_terms = [t for t in (other_concept_terms + other_item_terms) if t != correct_ans and t.lower() not in correct_ans.lower() and correct_ans.lower() not in t.lower()]
                fallback_concepts = banks.get("concepts", cs_banks["concepts"])
                opts = build_smart_options(correct_ans, other_terms, fallback_concepts)
                questions.append({
                    "prompt": prompt,
                    "options": opts,
                    "answer": correct_ans,
                    "explanation": f"According to the study material, {correct_ans} {cs['verb']} {cs['predicate']}.",
                    "skill": "Conceptual Identification",
                    "difficulty": "Medium"
                })

    return questions[:max_questions]



GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
GEMINI_SOURCE_LIMIT = 16_000
GEMINI_CANDIDATE_MODELS = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-2.5-pro", "gemini-1.5-flash"]


def _parse_gemini_json(text):
    """Extract and parse valid JSON even if wrapped in markdown code fences."""
    text = text.strip()
    if "```" in text:
        match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
        if match:
            text = match.group(1)
        else:
            text = re.sub(r"^```(?:json)?\s*", "", text)
            text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text.strip())
    except Exception:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def _valid_study_pack(pack):
    if not isinstance(pack, dict):
        return None
    raw_summary = pack.get("summary")
    raw_questions = pack.get("questions")
    if not isinstance(raw_summary, list) or not isinstance(raw_questions, list):
        return None
    summary = [point.strip() for point in raw_summary if isinstance(point, str) and point.strip()]
    questions = []
    for item in raw_questions:
        if not isinstance(item, dict):
            continue
        options = item.get("options")
        if not isinstance(options, list):
            continue
        options = [str(option).strip() for option in options if str(option).strip()]
        if len(options) != 4 or len(set(options)) != 4:
            continue
        answer = str(item.get("answer", "")).strip()
        prompt = str(item.get("prompt", "")).strip()
        explanation = str(item.get("explanation", "")).strip()
        if not prompt:
            continue
        matching_opt = next((opt for opt in options if opt.lower() == answer.lower()), None)
        if not matching_opt:
            continue
        answer = matching_opt
        difficulty = str(item.get("difficulty", "Easy")).capitalize()
        if difficulty not in {"Easy", "Medium", "Hard"}:
            difficulty = "Easy"
        skill = str(item.get("skill", "Key concept")).strip() or "Key concept"
        if not explanation:
            explanation = f"Review the related source point for a deeper understanding."
        
        shuffled_options = list(options)
        random.shuffle(shuffled_options)
        
        questions.append({
            "prompt": prompt,
            "options": shuffled_options,
            "answer": answer,
            "explanation": explanation,
            "skill": skill,
            "difficulty": difficulty,
        })
    if len(summary) < 1 or len(questions) < 1:
        return None
    return {"summary": summary[:5], "questions": questions[:30]}


def study_pack_prompt(content):
    source = content.strip()[:GEMINI_SOURCE_LIMIT]
    return f"""You are an elite university professor and gold-standard examination designer (e.g. USMLE / ETS / GRE / AP standard).
Your task is to create a high-quality, rigorous academic study pack strictly grounded in the SOURCE material below.

Strict MCQ Quality & Distractor Standards:
1. Professional Question Stems: Every question MUST be formulated as a complete, formal, grammatically polished academic question (e.g., "What is the primary physiological mechanism of...", "Which algorithmic property guarantees...", "Under what operational condition does...", "How does component X interact with Y?").
2. Homogeneous Option Length & Grammar: All 4 options (A, B, C, D) MUST have approximately equal length (within ±15% character count) and share identical grammatical construction (e.g., all four options must be complete verb phrases, or all four must be parallel noun phrases). No single option should stand out by being noticeably longer, more nuanced, or more detailed than the others.
3. High-Plausibility Academic Distractors: Distractors MUST represent credible, realistic academic alternatives and domain-relevant misconceptions. NEVER use silly, generic, or obviously incorrect options. All 4 choices must look equally plausible and sophisticated to someone who has not mastered the material.
4. Mutually Exclusive & Strictly Distinct: All 4 options MUST be distinct and non-overlapping. NEVER use "All of the above", "None of the above", or trivial variations.
5. Difficulty Progression:
   - Easy: Direct conceptual identification, foundational definitions, and primary components.
   - Medium: Process flow, functional roles, interactions, and mechanisms.
   - Hard: Comparative trade-offs, edge-case evaluations, failure modes, and architectural implications.
6. Pedagogical Explanations: Provide a concise, clear rationale explaining WHY the correct option is true and citing the underlying principle from the text.
7. Question Bank Volume: Generate between 20 to 30 high-yield, comprehensive questions covering all key aspects of the source.
8. Summary: 3 to 5 clear, insightful executive summary bullet points covering key takeaways.

Return ONLY a valid JSON object with the following schema:
{{
  "summary": ["Executive study point 1", "Executive study point 2", "Executive study point 3"],
  "questions": [
    {{
      "prompt": "Full professional academic question stem here?",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "answer": "Exact matching string from the 4 options",
      "explanation": "Concise pedagogical rationale explaining why this option is correct based on the source.",
      "skill": "Conceptual Identification / Architectural Analysis / Process Evaluation",
      "difficulty": "Easy" or "Medium" or "Hard"
    }}
  ]
}}

SOURCE MATERIAL:
\"\"\"
{source}
\"\"\""""


def gemini_study_pack(content):
    """Generate a predictable, source-grounded learning pack with multi-model fallback."""
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key or not content.strip():
        return None
    preferred_model = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash").strip()
    if preferred_model == "gemini-3.5-flash" or not preferred_model:
        preferred_model = "gemini-2.5-flash"
    
    models_to_try = [preferred_model] + [m for m in GEMINI_CANDIDATE_MODELS if m != preferred_model]
    prompt = study_pack_prompt(content)
    payload = json.dumps({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "responseMimeType": "application/json",
            "temperature": 0.3,
        },
    }).encode("utf-8")

    for model_name in models_to_try:
        endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent"
        request = urlrequest.Request(endpoint, data=payload, headers={
            "Content-Type": "application/json",
            "x-goog-api-key": api_key,
        }, method="POST")
        try:
            with urlrequest.urlopen(request, timeout=60) as response:
                body = json.loads(response.read().decode("utf-8"))
            text = "".join(part.get("text", "") for part in body["candidates"][0]["content"]["parts"])
            parsed = _parse_gemini_json(text)
            study_pack = _valid_study_pack(parsed)
            if study_pack:
                logging.info(f"Successfully generated study pack with model {model_name}.")
                return study_pack
        except Exception as error:
            logging.warning(f"Gemini generation with {model_name} failed ({error}). Trying next fallback model...")
            continue

    logging.warning("All Gemini candidate models were exhausted or unavailable.")
    return None


def local_study_pack(content):
    return {
        "summary": summarize(content),
        "questions": make_quiz(content),
    }


def process_document(document_id):
    with app.app_context():
        db = get_db()
        document = db.execute("SELECT content FROM documents WHERE id = ?", (document_id,)).fetchone()
        if not document:
            return
        db.execute("UPDATE documents SET processing_status = 'processing' WHERE id = ?", (document_id,))
        db.commit()
        gemini_pack = gemini_study_pack(document["content"])
        study_pack = gemini_pack or local_study_pack(document["content"])
        summary_points = study_pack["summary"]
        questions = study_pack["questions"]
        db.execute("DELETE FROM study_summaries WHERE document_id = ?", (document_id,))
        db.execute("DELETE FROM quiz_questions WHERE document_id = ?", (document_id,))
        db.executemany(
            "INSERT INTO study_summaries (document_id, position, sentence) VALUES (?, ?, ?)",
            [(document_id, position, sentence) for position, sentence in enumerate(summary_points)],
        )
        db.executemany(
            "INSERT INTO quiz_questions (document_id, position, prompt, options_json, answer) VALUES (?, ?, ?, ?, ?)",
            [(document_id, position, question["prompt"], json.dumps({key: value for key, value in question.items() if key != "prompt"}), question["answer"]) for position, question in enumerate(questions)],
        )
        db.execute("UPDATE documents SET processing_status = 'ready' WHERE id = ?", (document_id,))
        db.commit()


def queue_document_processing(document_id):
    db = get_db()
    db.execute("UPDATE documents SET processing_status = 'queued' WHERE id = ?", (document_id,))
    db.commit()
    if scheduler.running:
        scheduler.add_job(process_document, args=[document_id], id=f"document_{document_id}", replace_existing=True)
    else:
        process_document(document_id)


def stored_summary(document_id):
    return [row["sentence"] for row in get_db().execute("SELECT sentence FROM study_summaries WHERE document_id = ? ORDER BY position", (document_id,))]


def stored_questions(document_id):
    questions = []
    for row in get_db().execute("SELECT prompt, options_json, answer FROM quiz_questions WHERE document_id = ? ORDER BY position", (document_id,)):
        try:
            stored_data = json.loads(row["options_json"])
        except (TypeError, ValueError):
            continue
        # Earlier releases stored only the option list; retain access to those quizzes after the AI upgrade.
        metadata = stored_data if isinstance(stored_data, dict) else {"options": stored_data}
        question = {"prompt": row["prompt"], "answer": row["answer"]}
        question.update(metadata)
        question.setdefault("difficulty", "Easy")
        question.setdefault("skill", "Key concept")
        question.setdefault("explanation", "Review the related source point for a deeper understanding.")
        questions.append(question)
    return questions


QUIZ_FOCUSES = {
    "balanced": ("Balanced", ["Easy", "Medium", "Hard"]),
    "easy": ("Easy focus", ["Easy", "Medium", "Hard"]),
    "medium": ("Medium focus", ["Medium", "Easy", "Hard"]),
    "hard": ("Hard focus", ["Hard", "Medium", "Easy"]),
}


def select_quiz_questions(question_bank, focus, count, seed=None):
    """Select from the shared bank; uses optional seed for fair randomized sampling across attempts."""
    focus = focus if focus in QUIZ_FOCUSES else "balanced"
    count = max(1, min(count, 30, len(question_bank)))
    rng = random.Random(seed) if seed is not None else random.Random()
    buckets = {difficulty: [] for difficulty in ("Easy", "Medium", "Hard")}
    for question in question_bank:
        q_copy = dict(question)
        buckets.get(q_copy.get("difficulty", "Easy"), buckets["Easy"]).append(q_copy)
    for diff in buckets:
        rng.shuffle(buckets[diff])
    if focus == "balanced":
        selected = []
        while len(selected) < count and any(buckets.values()):
            for difficulty in QUIZ_FOCUSES[focus][1]:
                if buckets[difficulty] and len(selected) < count:
                    selected.append(buckets[difficulty].pop(0))
    else:
        selected = []
        for difficulty in QUIZ_FOCUSES[focus][1]:
            selected.extend(buckets[difficulty])
        selected = selected[:count]
    # Reshuffle options per question using the seed so positions vary
    for q in selected:
        if "options" in q and isinstance(q["options"], list):
            shuffled_opts = list(q["options"])
            rng.shuffle(shuffled_opts)
            q["options"] = shuffled_opts
    return selected


def quiz_count(value, available):
    try:
        requested = int(value)
    except (TypeError, ValueError):
        requested = min(10, available)
    return max(1, min(requested, 30, available)) if available else 0


def learner_performance(user_id, document_id=None):
    query = "SELECT COUNT(*) AS attempts, ROUND(AVG(score * 100.0 / total), 1) AS average_score, MAX(ROUND(score * 100.0 / total, 1)) AS best_score FROM quiz_attempts WHERE user_id = ?"
    parameters = [user_id]
    if document_id is not None:
        query += " AND document_id = ?"
        parameters.append(document_id)
    return get_db().execute(query, parameters).fetchone()


class DiffRowList(list):
    def __init__(self, rows, stats=None):
        super().__init__(rows)
        self.stats = stats or {}


def _html_escape(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _word_diff(old_line, new_line):
    """Perform intra-line word-level differential highlighting."""
    old_words = re.findall(r"\S+|\s+", old_line)
    new_words = re.findall(r"\S+|\s+", new_line)
    matcher = difflib.SequenceMatcher(None, old_words, new_words)
    old_parts = []
    new_parts = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        old_chunk = _html_escape("".join(old_words[i1:i2]))
        new_chunk = _html_escape("".join(new_words[j1:j2]))
        if tag == "equal":
            old_parts.append(old_chunk)
            new_parts.append(new_chunk)
        elif tag == "delete":
            old_parts.append(f'<mark class="diff-word-del font-semibold bg-rose-200 text-rose-950 px-1 rounded line-through">{old_chunk}</mark>')
        elif tag == "insert":
            new_parts.append(f'<mark class="diff-word-add font-semibold bg-emerald-200 text-emerald-950 px-1 rounded">{new_chunk}</mark>')
        elif tag == "replace":
            old_parts.append(f'<mark class="diff-word-del font-semibold bg-rose-200 text-rose-950 px-1 rounded line-through">{old_chunk}</mark>')
            new_parts.append(f'<mark class="diff-word-add font-semibold bg-emerald-200 text-emerald-950 px-1 rounded">{new_chunk}</mark>')
    return "".join(old_parts), "".join(new_parts)


def comparison_rows(original_content, version_content):
    original_lines = original_content.splitlines()
    version_lines = version_content.splitlines()
    rows = []
    matcher = difflib.SequenceMatcher(None, original_lines, version_lines)
    
    old_lineno = 1
    new_lineno = 1
    added_lines = 0
    removed_lines = 0
    modified_lines = 0
    unchanged_lines = 0

    for operation, old_start, old_end, new_start, new_end in matcher.get_opcodes():
        old_group = original_lines[old_start:old_end]
        new_group = version_lines[new_start:new_end]
        max_rows = max(len(old_group), len(new_group))
        
        for index in range(max_rows):
            has_old = index < len(old_group)
            has_new = index < len(new_group)
            old_line = old_group[index] if has_old else ""
            new_line = new_group[index] if has_new else ""
            
            curr_old_no = old_lineno if has_old else None
            curr_new_no = new_lineno if has_new else None
            
            if operation == "equal":
                old_state = new_state = "unchanged"
                row_type = "unchanged"
                old_html = _html_escape(old_line)
                new_html = _html_escape(new_line)
                unchanged_lines += 1
            elif operation == "delete":
                old_state, new_state = "removed", "empty"
                row_type = "removed"
                old_html = f'<mark class="diff-word-del font-semibold bg-rose-200 text-rose-950 px-1 rounded line-through">{_html_escape(old_line)}</mark>'
                new_html = ""
                removed_lines += 1
            elif operation == "insert":
                old_state, new_state = "empty", "added"
                row_type = "added"
                old_html = ""
                new_html = f'<mark class="diff-word-add font-semibold bg-emerald-200 text-emerald-950 px-1 rounded">{_html_escape(new_line)}</mark>'
                added_lines += 1
            else:  # replace
                old_state, new_state = "removed", "added"
                row_type = "modified"
                if has_old and has_new:
                    old_html, new_html = _word_diff(old_line, new_line)
                    modified_lines += 1
                elif has_old:
                    old_html = f'<mark class="diff-word-del font-semibold bg-rose-200 text-rose-950 px-1 rounded line-through">{_html_escape(old_line)}</mark>'
                    new_html = ""
                    removed_lines += 1
                else:
                    old_html = ""
                    new_html = f'<mark class="diff-word-add font-semibold bg-emerald-200 text-emerald-950 px-1 rounded">{_html_escape(new_line)}</mark>'
                    added_lines += 1
            
            if has_old:
                old_lineno += 1
            if has_new:
                new_lineno += 1
                
            rows.append({
                "old": old_line,
                "new": new_line,
                "old_state": old_state,
                "new_state": new_state,
                "old_num": curr_old_no,
                "new_num": curr_new_no,
                "type": row_type,
                "old_html": old_html,
                "new_html": new_html,
                "is_changed": row_type != "unchanged",
            })

    orig_words = len(re.findall(r"\w+", original_content))
    ver_words = len(re.findall(r"\w+", version_content))
    sim_ratio = round(difflib.SequenceMatcher(None, original_content, version_content).ratio() * 100, 1)
    
    stats = {
        "similarity_pct": sim_ratio,
        "orig_words": orig_words,
        "version_words": ver_words,
        "word_delta": ver_words - orig_words,
        "added_lines": added_lines,
        "removed_lines": removed_lines,
        "modified_lines": modified_lines,
        "unchanged_lines": unchanged_lines,
        "total_rows": len(rows),
    }
    return DiffRowList(rows, stats)


def inactivity_days(db):
    return int(db.execute("SELECT setting_value FROM system_settings WHERE setting_key = 'inactivity_days'").fetchone()[0])


def run_inactivity_watchdog():
    with app.app_context():
        db = get_db()
        cutoff = datetime.now(UTC) - timedelta(days=inactivity_days(db))
        documents = db.execute(
            "SELECT d.id, d.owner_id FROM documents d JOIN users u ON u.id = d.owner_id "
            "WHERE d.visibility = 'private' AND d.release_on_inactivity = 1 AND u.role = 'student' AND u.graduation_status = 'graduated' AND u.last_login < ?",
            (cutoff.isoformat(),),
        ).fetchall()
        if documents:
            db.executemany("UPDATE documents SET visibility = 'public' WHERE id = ?", [(document["id"],) for document in documents])
            db.executemany(
                "INSERT INTO activity_logs (user_id, document_id, event_type, detail, created_at) VALUES (?, ?, 'watchdog_release', 'Released by inactivity policy', ?)",
                [(document["owner_id"], document["id"], now_iso()) for document in documents],
            )
        db.execute("UPDATE system_settings SET setting_value = ? WHERE setting_key = 'last_watchdog_run'", (now_iso(),))
        db.commit()
        return len(documents)


def start_scheduler():
    if scheduler.running:
        return
    scheduler.add_job(run_inactivity_watchdog, "interval", days=1, id="inactivity_watchdog", replace_existing=True)
    scheduler.start()
    atexit.register(lambda: scheduler.shutdown(wait=False))


@app.route("/")
def index():
    return redirect(url_for("dashboard")) if g.user else redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        user = get_db().execute("SELECT * FROM users WHERE email = ?", (request.form["email"].strip().lower(),)).fetchone()
        if user and user["password_hash"] and check_password_hash(user["password_hash"], request.form["password"]) and user["is_verified"]:
            session.clear()
            session["user_id"] = user["id"]
            db = get_db()
            db.execute("UPDATE users SET last_login = ? WHERE id = ?", (now_iso(), user["id"]))
            db.execute("INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'login', 'Successful login', ?)", (user["id"], now_iso()))
            db.commit()
            return redirect(url_for("dashboard"))
        if user and not user["is_verified"]:
            flash("Your account is waiting for administrator verification.", "error")
        else:
            flash("Invalid email or password.", "error")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if g.user:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        name = request.form["name"].strip()
        reg_no = request.form.get("reg_no", "").strip()
        email = request.form["email"].strip().lower()
        password = request.form["password"]
        role = request.form["role"]
        if not name or not email or len(password) < 8 or role not in {"student", "staff"}:
            flash("Use a name, valid email, 8-character password, and student or staff role.", "error")
        else:
            try:
                db = get_db()
                cursor = db.execute(
                    "INSERT INTO users (name, reg_no, email, password, password_hash, role, department, semester, last_login, is_verified) VALUES (?, ?, ?, '', ?, ?, ?, ?, ?, 0)",
                    (name, reg_no, email, generate_password_hash(password), role, request.form["department"].strip(), request.form["semester"].strip(), now_iso()),
                )
                db.execute("INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'registration', 'Verification requested', ?)", (cursor.lastrowid, now_iso()))
                db.commit()
                flash("Registration submitted. An administrator must verify your account before you sign in.", "success")
                return redirect(url_for("login"))
            except sqlite3.IntegrityError:
                flash("An account already uses that email address.", "error")
    return render_template("register.html")


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def get_staff_accessible_doc_ids(user):
    db = get_db()
    u = dict(user) if user else {}
    if u.get("role") == "admin":
        return [d["id"] for d in db.execute("SELECT id FROM documents").fetchall()]
    doc_ids = [d["id"] for d in db.execute("SELECT id FROM documents WHERE owner_id = ?", (u.get("id"),)).fetchall()]
    if not doc_ids:
        doc_ids = [d["id"] for d in db.execute("SELECT id FROM documents WHERE department = ? OR visibility = 'public'", (u.get("department"),)).fetchall()]
    return doc_ids


def format_report_datetime(iso_str):
    if not iso_str:
        return "N/A"
    try:
        dt = datetime.fromisoformat(str(iso_str).replace("Z", "+00:00"))
        return dt.strftime("%Y-%m-%d %H:%M")
    except Exception:
        return str(iso_str)[:16].replace("T", " ")


def get_performance_grade(score_pct):
    if score_pct is None:
        return "Unattempted"
    try:
        score = float(score_pct)
    except (ValueError, TypeError):
        return "Unattempted"
    if score >= 85:
        return "Distinction (85-100%)"
    if score >= 70:
        return "Proficient (70-84%)"
    if score >= 50:
        return "Developing (50-69%)"
    return "Needs Remediation (<50%)"


def get_staff_comprehensive_metrics(user):
    db = get_db()
    staff_doc_ids = get_staff_accessible_doc_ids(user)
    placeholders = ",".join("?" for _ in staff_doc_ids) if staff_doc_ids else "?"
    params = staff_doc_ids if staff_doc_ids else [-1]

    # Monitored Notes Performance
    staff_notes = db.execute(f"""
        SELECT d.id, d.title, d.subject, d.department, d.semester, d.visibility, d.created_at, d.owner_id, u.name AS owner_name,
               (SELECT COUNT(*) FROM activity_logs a WHERE a.document_id = d.id AND a.event_type = 'document_view') AS views_count,
               (SELECT COUNT(*) FROM quiz_attempts q WHERE q.document_id = d.id) AS quiz_attempts_count,
               (SELECT COUNT(DISTINCT q.user_id) FROM quiz_attempts q WHERE q.document_id = d.id) AS unique_students_count,
               (SELECT ROUND(AVG(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END), 1) FROM quiz_attempts q WHERE q.document_id = d.id) AS avg_mastery,
               (SELECT MAX(ROUND(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END, 1)) FROM quiz_attempts q WHERE q.document_id = d.id) AS top_score
        FROM documents d
        JOIN users u ON u.id = d.owner_id
        WHERE d.id IN ({placeholders})
        ORDER BY d.created_at DESC
    """, params).fetchall()

    # Learners Performance Roster (excludes current faculty member from their own student list)
    user_filter_params = params + [user["id"]]
    staff_learners = db.execute(f"""
        SELECT u.id, u.name, u.email, u.department, u.semester,
               COUNT(q.id) AS total_attempts,
               ROUND(AVG(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END), 1) AS avg_score,
               MAX(ROUND(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END, 1)) AS max_score,
               MAX(q.created_at) AS last_active
        FROM users u
        JOIN quiz_attempts q ON q.user_id = u.id
        WHERE q.document_id IN ({placeholders}) AND u.id != ?
        GROUP BY u.id
        ORDER BY avg_score DESC, total_attempts DESC
    """, user_filter_params).fetchall()

    # Knowledge gaps (Notes with avg < 75%)
    staff_gaps = [n for n in staff_notes if n["avg_mastery"] is not None and n["avg_mastery"] < 75]

    # Grade distribution across all individual attempts
    all_scores = [r[0] for r in db.execute(f"""
        SELECT CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END
        FROM quiz_attempts q
        JOIN users u ON u.id = q.user_id
        WHERE q.document_id IN ({placeholders}) AND u.id != ?
    """, user_filter_params).fetchall() if r[0] is not None]

    total_views = sum(n["views_count"] for n in staff_notes)
    total_attempts = len(all_scores)
    overall_avg = round(sum(all_scores) / max(1, len(all_scores)), 1) if all_scores else None

    staff_distribution = {
        "A": sum(1 for s in all_scores if s >= 85),
        "B": sum(1 for s in all_scores if 70 <= s < 85),
        "C": sum(1 for s in all_scores if 50 <= s < 70),
        "D": sum(1 for s in all_scores if s < 50),
        "total": len(all_scores),
    }

    at_risk_learners = [l for l in staff_learners if l["avg_score"] is not None and l["avg_score"] < 50]

    overview = {
        "attempts": total_attempts,
        "learners": len(staff_learners),
        "average_score": overall_avg,
        "materials": len(staff_notes),
    }

    return {
        "staff_notes": staff_notes,
        "staff_learners": staff_learners,
        "staff_gaps": staff_gaps,
        "all_scores": all_scores,
        "total_views": total_views,
        "total_attempts": total_attempts,
        "overall_avg": overall_avg,
        "staff_distribution": staff_distribution,
        "at_risk_learners": at_risk_learners,
        "interventions": staff_gaps,
        "overview": overview,
    }


@app.route("/dashboard")
@login_required()
def dashboard():
    db = get_db()
    accessible_materials = accessible_documents()
    recent_documents = accessible_materials[:8]
    recent_attempts = []
    subject_stats = []
    pending_materials = []
    needs_review = []
    perfect_count = 0

    staff_notes = []
    staff_learners = []
    staff_gaps = []
    staff_distribution = {"A": 0, "B": 0, "C": 0, "D": 0, "total": 0}
    pending_profile_request = None

    if g.user["role"] == "student":
        pending_profile_request = db.execute(
            "SELECT * FROM profile_update_requests WHERE user_id = ? AND status = 'pending' ORDER BY created_at DESC LIMIT 1",
            (g.user["id"],)
        ).fetchone()
        overall_performance = learner_performance(g.user["id"])
        attempts_count = db.execute("SELECT COUNT(*) FROM quiz_attempts WHERE user_id = ?", (g.user["id"],)).fetchone()[0]
        perfect_count = db.execute("SELECT COUNT(*) FROM quiz_attempts WHERE user_id = ? AND score = total AND total > 0", (g.user["id"],)).fetchone()[0]
        
        role_metrics = [
            ("Enrolled Materials", len(accessible_materials)),
            ("Quiz Attempts", attempts_count),
            ("Average Mastery", f"{overall_performance['average_score'] or 0}%"),
            ("Best Score", f"{overall_performance['best_score'] or 0}%"),
        ]
        workspace_label = f"Student Academic Workspace · {g.user['department']} (Sem {g.user['semester']})"
        
        # Subject Mastery Breakdown
        subject_stats = db.execute(
            "SELECT d.subject, COUNT(q.id) as attempts, ROUND(AVG(q.score * 100.0 / q.total), 1) as avg_score, MAX(ROUND(q.score * 100.0 / q.total), 1) as max_score "
            "FROM quiz_attempts q JOIN documents d ON d.id = q.document_id "
            "WHERE q.user_id = ? GROUP BY d.subject ORDER BY avg_score DESC",
            (g.user["id"],)
        ).fetchall()

        # Upcoming Work & Action Items:
        accessible_doc_ids = {doc["id"] for doc in accessible_materials}
        attempted_doc_ids = [row[0] for row in db.execute("SELECT DISTINCT document_id FROM quiz_attempts WHERE user_id = ?", (g.user["id"],)).fetchall() if row[0] in accessible_doc_ids]
        pending_materials = [doc for doc in accessible_materials if doc["id"] not in attempted_doc_ids][:4]

        # Revision recommendations (last score < 75%)
        for doc_id in attempted_doc_ids:
            latest = db.execute(
                "SELECT q.*, d.title, d.subject, ROUND(q.score * 100.0 / q.total, 1) as pct "
                "FROM quiz_attempts q JOIN documents d ON d.id = q.document_id "
                "WHERE q.user_id = ? AND q.document_id = ? ORDER BY q.created_at DESC LIMIT 1",
                (g.user["id"], doc_id)
            ).fetchone()
            if latest and latest["pct"] < 75:
                needs_review.append(latest)

        all_recent = db.execute(
            "SELECT q.*, d.title, d.subject, ROUND(q.score * 100.0 / q.total, 1) AS mastery_percent "
            "FROM quiz_attempts q JOIN documents d ON d.id = q.document_id WHERE q.user_id = ? "
            "ORDER BY q.created_at DESC",
            (g.user["id"],),
        ).fetchall()
        recent_attempts = [att for att in all_recent if att["document_id"] in accessible_doc_ids][:6]

    elif g.user["role"] == "staff":
        metrics = get_staff_comprehensive_metrics(g.user)
        staff_notes = metrics["staff_notes"]
        staff_learners = metrics["staff_learners"]
        staff_gaps = metrics["staff_gaps"]
        staff_distribution = metrics["staff_distribution"]
        total_views = metrics["total_views"]
        total_attempts = metrics["total_attempts"]
        overall_avg = metrics["overall_avg"]

        role_metrics = [
            ("Monitored Notes", len(staff_notes)),
            ("Total Student Views", total_views),
            ("Class Assessment Attempts", total_attempts),
            ("Class Avg Mastery", f"{overall_avg}%" if overall_avg is not None else "—"),
        ]
        workspace_label = f"Teaching Intelligence Hub · {g.user['department']}"
    else:
        role_metrics = [
            ("Registered users", db.execute("SELECT COUNT(*) FROM users").fetchone()[0]),
            ("Repository materials", db.execute("SELECT COUNT(*) FROM documents WHERE visibility = 'public'").fetchone()[0]),
            ("Recorded events", db.execute("SELECT COUNT(*) FROM activity_logs").fetchone()[0]),
        ]
        workspace_label = "System overview"

    return render_template(
        "dashboard.html",
        documents=recent_documents,
        role_metrics=role_metrics,
        workspace_label=workspace_label,
        recent_attempts=recent_attempts,
        subject_stats=subject_stats,
        pending_materials=pending_materials,
        needs_review=needs_review,
        perfect_count=perfect_count,
        staff_notes=staff_notes,
        staff_learners=staff_learners,
        staff_gaps=staff_gaps,
        staff_distribution=staff_distribution,
        pending_profile_request=pending_profile_request,
    )


def build_staff_excel_workbook(user, report_type="master"):
    if not openpyxl:
        return None

    wb = openpyxl.Workbook()
    # Remove default placeholder sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    metrics = get_staff_comprehensive_metrics(user)
    staff_notes = metrics["staff_notes"]
    staff_learners = metrics["staff_learners"]
    overview = metrics["overview"]
    total_views = metrics["total_views"]

    u = dict(user) if user else {}
    faculty_name = u.get("name") or "Faculty Instructor"
    department = u.get("department") or "Academic Department"
    gen_time_str = datetime.now().strftime("%B %d, %Y at %I:%M %p")

    # Reusable styling objects
    thin_border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    col_header_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
    col_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    banner_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    banner_font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")

    sub_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    kpi_label_font = Font(name="Segoe UI", size=8.5, bold=True, color="4338CA")
    kpi_value_font = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    kpi_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    row_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    data_font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    badge_styles = {
        "Distinction": (PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="166534")),
        "Proficient": (PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="3730A3")),
        "Developing": (PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="92400E")),
        "Remediation": (PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"), Font(name="Segoe UI", size=9.5, bold=True, color="991B1B")),
        "Pending": (PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"), Font(name="Segoe UI", size=9.5, italic=True, color="64748B")),
    }

    # 1. Student Gradebook Sheet
    if report_type in {"master", "summary", "gradebook"}:
        ws1 = wb.create_sheet(title="Student Gradebook")
        ws1.views.sheetView[0].showGridLines = True

        ws1.merge_cells("A1:K1")
        ws1["A1"] = "OmniStudy · Student Competency & Gradebook Ledger"
        ws1["A1"].font = banner_font
        ws1["A1"].fill = banner_fill
        ws1["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws1.row_dimensions[1].height = 32

        ws1.merge_cells("A2:K2")
        ws1["A2"] = f"Department: {department}  |  Faculty: {faculty_name}  |  Generated: {gen_time_str}"
        ws1["A2"].font = sub_font
        ws1["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws1.row_dimensions[2].height = 20

        # KPI blocks (Rows 4-5)
        kpis = [
            ("A", "C", "UNIQUE ASSESSED STUDENTS", str(overview["learners"])),
            ("D", "F", "TOTAL QUIZ ATTEMPTS", str(overview["attempts"])),
            ("G", "I", "COHORT AVERAGE MASTERY", f"{overview['average_score']}%" if overview["average_score"] is not None else "—"),
            ("J", "K", "MONITORED NOTES", str(overview["materials"])),
        ]
        ws1.row_dimensions[4].height = 16
        ws1.row_dimensions[5].height = 26
        for start_c, end_c, label, val in kpis:
            start_num = openpyxl.utils.column_index_from_string(start_c)
            end_num = openpyxl.utils.column_index_from_string(end_c)
            ws1.merge_cells(f"{start_c}4:{end_c}4")
            ws1.merge_cells(f"{start_c}5:{end_c}5")
            ws1[f"{start_c}4"] = label
            ws1[f"{start_c}4"].font = kpi_label_font
            ws1[f"{start_c}4"].alignment = align_center
            ws1[f"{start_c}5"] = val
            ws1[f"{start_c}5"].font = kpi_value_font
            ws1[f"{start_c}5"].alignment = align_center
            for r in [4, 5]:
                for c in range(start_num, end_num + 1):
                    cell = ws1.cell(row=r, column=c)
                    cell.fill = kpi_fill
                    cell.border = cell_border

        headers1 = [
            "Student ID", "Student Name", "Email Address", "Department", "Semester",
            "Quiz Attempts", "Average Mastery", "Best Score", "Grade Band", "Academic Standing", "Last Active Date"
        ]
        h_row1 = 7
        ws1.row_dimensions[h_row1].height = 26
        for col_idx, h in enumerate(headers1, start=1):
            cell = ws1.cell(row=h_row1, column=col_idx, value=h)
            cell.font = col_header_font
            cell.fill = col_header_fill
            cell.alignment = align_center
            cell.border = cell_border

        cur_r = 8
        for l in staff_learners:
            ws1.row_dimensions[cur_r].height = 22
            row_fill = row_alt_fill if cur_r % 2 == 0 else row_white_fill
            avg_num = (l["avg_score"] or 0) / 100.0 if l["avg_score"] is not None else 0
            best_num = (l["max_score"] or 0) / 100.0 if l["max_score"] is not None else 0

            if l["avg_score"] is None:
                standing = "No Attempts Logged"
                b_key = "Pending"
            elif l["avg_score"] < 50:
                standing = "Intervention Recommended"
                b_key = "Remediation"
            elif l["avg_score"] >= 85:
                standing = "High Mastery / Honors"
                b_key = "Distinction"
            elif l["avg_score"] >= 70:
                standing = "Proficient / Solid"
                b_key = "Proficient"
            else:
                standing = "Developing Baseline"
                b_key = "Developing"

            grade_label = get_performance_grade(l["avg_score"])
            b_fill, b_font = badge_styles[b_key]

            row_cells = [
                (l["id"], align_center, data_font),
                (l["name"], align_left, data_font_bold),
                (l["email"], align_left, data_font),
                (l["department"], align_center, data_font),
                (f"Sem {l['semester']}", align_center, data_font),
                (l["total_attempts"], align_center, data_font),
                (avg_num, align_center, data_font_bold),
                (best_num, align_center, data_font),
                (grade_label, align_center, b_font),
                (standing, align_center, data_font),
                (format_report_datetime(l["last_active"]), align_center, data_font),
            ]

            for c_idx, (val, align, font) in enumerate(row_cells, start=1):
                cell = ws1.cell(row=cur_r, column=c_idx, value=val)
                cell.font = font
                cell.alignment = align
                cell.border = cell_border
                if c_idx == 9:
                    cell.fill = b_fill
                else:
                    cell.fill = row_fill
                if c_idx in (7, 8):
                    cell.number_format = "0.0%"

            cur_r += 1

        last_r1 = max(cur_r - 1, h_row1)
        ws1.auto_filter.ref = f"A{h_row1}:K{last_r1}"
        ws1.freeze_panes = "A8"

        for col in ws1.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 4, 13)

    # 2. Material Outcomes Sheet
    if report_type in {"master", "materials"}:
        ws2 = wb.create_sheet(title="Material Outcomes")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:K1")
        ws2["A1"] = "OmniStudy · Curriculum Material Assessment Outcomes Audit"
        ws2["A1"].font = banner_font
        ws2["A1"].fill = banner_fill
        ws2["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws2.row_dimensions[1].height = 32

        ws2.merge_cells("A2:K2")
        ws2["A2"] = f"Department: {department}  |  Monitored Notes: {len(staff_notes)}  |  Generated: {gen_time_str}"
        ws2["A2"].font = sub_font
        ws2["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws2.row_dimensions[2].height = 20

        m_headers = [
            "Material ID", "Lecture Note Title", "Subject", "Department", "Semester",
            "Visibility", "Readership Views", "Unique Students", "Quiz Attempts",
            "Average Mastery", "Curriculum Health Status"
        ]
        h_row2 = 4
        ws2.row_dimensions[h_row2].height = 26
        for col_idx, h in enumerate(m_headers, start=1):
            cell = ws2.cell(row=h_row2, column=col_idx, value=h)
            cell.font = col_header_font
            cell.fill = col_header_fill
            cell.alignment = align_center
            cell.border = cell_border

        cur_r2 = 5
        for m in staff_notes:
            ws2.row_dimensions[cur_r2].height = 22
            row_fill = row_alt_fill if cur_r2 % 2 == 0 else row_white_fill
            avg_num = (m["avg_mastery"] or 0) / 100.0 if m["avg_mastery"] is not None else 0

            if m["avg_mastery"] is None:
                health = "Pending Student Attempts"
                b_key = "Pending"
            elif m["avg_mastery"] >= 85:
                health = "Exemplary (>=85%)"
                b_key = "Distinction"
            elif m["avg_mastery"] >= 70:
                health = "Proficient (70-84%)"
                b_key = "Proficient"
            elif m["avg_mastery"] >= 50:
                health = "Developing (50-69%)"
                b_key = "Developing"
            else:
                health = "Critical Review Needed (<50%)"
                b_key = "Remediation"

            b_fill, b_font = badge_styles[b_key]

            m_row = [
                (m["id"], align_center, data_font),
                (m["title"], align_left, data_font_bold),
                (m["subject"], align_left, data_font),
                (m["department"], align_center, data_font),
                (f"Sem {m['semester']}", align_center, data_font),
                ((m["visibility"] or "").title(), align_center, data_font),
                (m["views_count"], align_center, data_font),
                (m["unique_students_count"], align_center, data_font),
                (m["quiz_attempts_count"], align_center, data_font),
                (avg_num if m["avg_mastery"] is not None else "N/A", align_center, data_font_bold),
                (health, align_center, b_font),
            ]

            for c_idx, (val, align, font) in enumerate(m_row, start=1):
                cell = ws2.cell(row=cur_r2, column=c_idx, value=val)
                cell.font = font
                cell.alignment = align
                cell.border = cell_border
                if c_idx == 11:
                    cell.fill = b_fill
                else:
                    cell.fill = row_fill
                if c_idx == 10 and m["avg_mastery"] is not None:
                    cell.number_format = "0.0%"

            cur_r2 += 1

        last_r2 = max(cur_r2 - 1, h_row2)
        ws2.auto_filter.ref = f"A{h_row2}:K{last_r2}"
        ws2.freeze_panes = "A5"

        for col in ws2.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 3. Quiz Attempts Ledger Sheet
    if report_type in {"master", "attempts"}:
        ws3 = wb.create_sheet(title="Quiz Attempts Ledger")
        ws3.views.sheetView[0].showGridLines = True

        db = get_db()
        staff_doc_ids = get_staff_accessible_doc_ids(user)
        placeholders = ",".join("?" for _ in staff_doc_ids) if staff_doc_ids else "?"
        params = staff_doc_ids if staff_doc_ids else [-1]
        user_filter_params = params + [u.get("id")]

        detailed_attempts = db.execute(f"""
            SELECT q.id, u.name AS student_name, u.email AS student_email, u.department AS student_department, u.semester AS student_semester,
                   d.title AS document_title, d.subject,
                   q.score, q.total,
                   ROUND(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END, 1) AS score_pct,
                   q.created_at
            FROM quiz_attempts q
            JOIN users u ON u.id = q.user_id
            JOIN documents d ON d.id = q.document_id
            WHERE q.document_id IN ({placeholders}) AND u.id != ?
            ORDER BY q.created_at DESC
        """, user_filter_params).fetchall()

        ws3.merge_cells("A1:K1")
        ws3["A1"] = "OmniStudy · Detailed Student Quiz Submissions Ledger"
        ws3["A1"].font = banner_font
        ws3["A1"].fill = banner_fill
        ws3["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws3.row_dimensions[1].height = 32

        ws3.merge_cells("A2:K2")
        ws3["A2"] = f"Total Attempt Records: {len(detailed_attempts)}  |  Faculty: {faculty_name}  |  Generated: {gen_time_str}"
        ws3["A2"].font = sub_font
        ws3["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws3.row_dimensions[2].height = 20

        att_headers = [
            "Attempt ID", "Student Name", "Email Address", "Department", "Semester",
            "Material Title", "Subject", "Score Earned", "Total Questions", "Score %", "Grade Band"
        ]
        h_row3 = 4
        ws3.row_dimensions[h_row3].height = 26
        for col_idx, h in enumerate(att_headers, start=1):
            cell = ws3.cell(row=h_row3, column=col_idx, value=h)
            cell.font = col_header_font
            cell.fill = col_header_fill
            cell.alignment = align_center
            cell.border = cell_border

        cur_r3 = 5
        for a in detailed_attempts:
            ws3.row_dimensions[cur_r3].height = 22
            row_fill = row_alt_fill if cur_r3 % 2 == 0 else row_white_fill
            pct_val = (float(a["score_pct"]) if a["score_pct"] is not None else 0) / 100.0
            grade_str = get_performance_grade(a["score_pct"])

            if pct_val >= 0.85:
                b_key = "Distinction"
            elif pct_val >= 0.70:
                b_key = "Proficient"
            elif pct_val >= 0.50:
                b_key = "Developing"
            else:
                b_key = "Remediation"

            b_fill, b_font = badge_styles[b_key]

            att_row = [
                (a["id"], align_center, data_font),
                (a["student_name"], align_left, data_font_bold),
                (a["student_email"], align_left, data_font),
                (a["student_department"], align_center, data_font),
                (f"Sem {a['student_semester']}", align_center, data_font),
                (a["document_title"], align_left, data_font),
                (a["subject"], align_left, data_font),
                (a["score"], align_center, data_font),
                (a["total"], align_center, data_font),
                (pct_val, align_center, data_font_bold),
                (grade_str, align_center, b_font),
            ]

            for c_idx, (val, align, font) in enumerate(att_row, start=1):
                cell = ws3.cell(row=cur_r3, column=c_idx, value=val)
                cell.font = font
                cell.alignment = align
                cell.border = cell_border
                if c_idx == 11:
                    cell.fill = b_fill
                else:
                    cell.fill = row_fill
                if c_idx == 10:
                    cell.number_format = "0.0%"

            cur_r3 += 1

        last_r3 = max(cur_r3 - 1, h_row3)
        ws3.auto_filter.ref = f"A{h_row3}:K{last_r3}"
        ws3.freeze_panes = "A5"

        for col in ws3.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws3.column_dimensions[col_letter].width = max(max_len + 4, 13)

    return wb


@app.route("/staff/export-excel")
@app.route("/staff/export/excel")
@app.route("/staff/export/excel/<string:report_type>")
@login_required({"staff", "admin"})
def export_staff_excel(report_type=None):
    rep_type = report_type or request.args.get("type", "master").lower()
    wb = build_staff_excel_workbook(g.user, rep_type)
    if not wb:
        flash("Excel engine is not initialized. Exporting as CSV instead.", "warning")
        return redirect(url_for("export_staff_roster", type=rep_type))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    u = dict(g.user) if g.user else {}
    dept_slug = re.sub(r"[^A-Za-z0-9_]+", "", (u.get("department") or "All").replace(" ", "_")) or "Dept"
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M")

    if rep_type in {"gradebook", "summary"}:
        filename = f"omnistudy_gradebook_{dept_slug}_{timestamp_str}.xlsx"
    elif rep_type == "materials":
        filename = f"omnistudy_materials_audit_{dept_slug}_{timestamp_str}.xlsx"
    elif rep_type == "attempts":
        filename = f"omnistudy_detailed_attempts_{dept_slug}_{timestamp_str}.xlsx"
    else:
        filename = f"omnistudy_teaching_dossier_{dept_slug}_{timestamp_str}.xlsx"

    response = Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/staff/export-roster")
@app.route("/staff/export/<string:report_type>")
@login_required({"staff", "admin"})
def export_staff_roster(report_type=None):
    export_type = report_type or request.args.get("type", "summary").lower()
    metrics = get_staff_comprehensive_metrics(g.user)
    staff_notes = metrics["staff_notes"]
    staff_learners = metrics["staff_learners"]
    dept_name = dict(g.user).get("department") or "All"
    dept_slug = re.sub(r"[^A-Za-z0-9_]+", "", dept_name.replace(" ", "_")) or "Dept"
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M")

    output = io.StringIO()
    # Write UTF-8 Byte Order Mark (BOM) for native Microsoft Excel compatibility on Windows
    output.write("\ufeff")
    writer = csv.writer(output)

    if export_type == "materials":
        writer.writerow([
            "Material ID",
            "Material Title",
            "Subject",
            "Department",
            "Semester",
            "Visibility",
            "Readership Views",
            "Unique Active Students",
            "Quiz Attempts",
            "Average Mastery (%)",
            "Top Score (%)",
            "Curriculum Health Status",
            "Publication Date",
        ])
        for m in staff_notes:
            avg_str = f"{m['avg_mastery']:.1f}" if m["avg_mastery"] is not None else "N/A"
            top_str = f"{m['top_score']:.1f}" if m["top_score"] is not None else "N/A"
            if m["avg_mastery"] is None:
                health = "Pending Student Attempts"
            elif m["avg_mastery"] >= 85:
                health = "Exemplary (>=85%)"
            elif m["avg_mastery"] >= 70:
                health = "Proficient (70-84%)"
            elif m["avg_mastery"] >= 50:
                health = "Developing (50-69%)"
            else:
                health = "Critical Review Needed (<50%)"

            writer.writerow([
                m["id"],
                m["title"],
                m["subject"],
                m["department"],
                m["semester"],
                (m["visibility"] or "").title(),
                m["views_count"],
                m["unique_students_count"],
                m["quiz_attempts_count"],
                avg_str,
                top_str,
                health,
                format_report_datetime(m["created_at"]),
            ])
        filename = f"omnistudy_materials_audit_{dept_slug}_{timestamp_str}.csv"

    elif export_type == "attempts":
        db = get_db()
        staff_doc_ids = get_staff_accessible_doc_ids(g.user)
        placeholders = ",".join("?" for _ in staff_doc_ids) if staff_doc_ids else "?"
        params = staff_doc_ids if staff_doc_ids else [-1]
        user_filter_params = params + [g.user["id"]]

        detailed_attempts = db.execute(f"""
            SELECT q.id, u.name AS student_name, u.email AS student_email, u.department AS student_department, u.semester AS student_semester,
                   d.title AS document_title, d.subject,
                   q.score, q.total,
                   ROUND(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END, 1) AS score_pct,
                   q.created_at
            FROM quiz_attempts q
            JOIN users u ON u.id = q.user_id
            JOIN documents d ON d.id = q.document_id
            WHERE q.document_id IN ({placeholders}) AND u.id != ?
            ORDER BY q.created_at DESC
        """, user_filter_params).fetchall()

        writer.writerow([
            "Attempt ID",
            "Student Name",
            "Email",
            "Department",
            "Semester",
            "Material Title",
            "Subject",
            "Score Earned",
            "Total Questions",
            "Mastery (%)",
            "Performance Grade",
            "Attempt Date & Time",
        ])
        for a in detailed_attempts:
            score_pct = float(a["score_pct"]) if a["score_pct"] is not None else 0.0
            writer.writerow([
                a["id"],
                a["student_name"],
                a["student_email"],
                a["student_department"],
                a["student_semester"],
                a["document_title"],
                a["subject"],
                a["score"],
                a["total"],
                f"{score_pct:.1f}",
                get_performance_grade(score_pct),
                format_report_datetime(a["created_at"]),
            ])
        filename = f"omnistudy_detailed_attempts_{dept_slug}_{timestamp_str}.csv"

    else:
        # Default / Gradebook Summary Roster
        writer.writerow([
            "Student ID",
            "Student Name",
            "Email",
            "Department",
            "Semester",
            "Quiz Attempts",
            "Average Mastery (%)",
            "Best Score (%)",
            "Performance Grade",
            "Academic Standing",
            "Last Active Date",
        ])
        for l in staff_learners:
            avg_score = l["avg_score"] if l["avg_score"] is not None else 0.0
            max_score = l["max_score"] if l["max_score"] is not None else 0.0
            if l["avg_score"] is None:
                standing = "No Attempts Logged"
            elif l["avg_score"] < 50:
                standing = "Academic Intervention Recommended"
            elif l["avg_score"] >= 85:
                standing = "Honors / High Mastery"
            else:
                standing = "Good Standing"

            writer.writerow([
                l["id"],
                l["name"],
                l["email"],
                l["department"],
                l["semester"],
                l["total_attempts"],
                f"{avg_score:.1f}",
                f"{max_score:.1f}",
                get_performance_grade(l["avg_score"]),
                standing,
                format_report_datetime(l["last_active"]),
            ])
        filename = f"omnistudy_gradebook_roster_{dept_slug}_{timestamp_str}.csv"

    response = Response(output.getvalue(), mimetype="text/csv; charset=utf-8")
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/staff/report")
@login_required({"staff", "admin"})
def staff_report():
    metrics = get_staff_comprehensive_metrics(g.user)
    gen_time = datetime.now()
    generation_date = gen_time.strftime("%B %d, %Y at %I:%M %p")
    generation_timestamp = gen_time.strftime("%Y%m%d%H%M%S")
    return render_template(
        "staff_report.html",
        staff_notes=metrics["staff_notes"],
        staff_learners=metrics["staff_learners"],
        total_views=metrics["total_views"],
        overview=metrics["overview"],
        staff_distribution=metrics["staff_distribution"],
        at_risk_learners=metrics["at_risk_learners"],
        interventions=metrics["interventions"],
        generation_date=generation_date,
        generation_timestamp=generation_timestamp,
        faculty_name=g.user["name"],
        department=dict(g.user).get("department") or "Academic Division",
    )


@app.route("/staff/report/download")
@login_required({"staff", "admin"})
def staff_report_download():
    metrics = get_staff_comprehensive_metrics(g.user)
    gen_time = datetime.now()
    generation_date = gen_time.strftime("%B %d, %Y at %I:%M %p")
    generation_timestamp = gen_time.strftime("%Y%m%d%H%M%S")
    dept_name = dict(g.user).get("department") or "Academic Division"
    dept_slug = re.sub(r"[^A-Za-z0-9_]+", "", dept_name.replace(" ", "_")) or "Dept"
    html_content = render_template(
        "staff_report.html",
        staff_notes=metrics["staff_notes"],
        staff_learners=metrics["staff_learners"],
        total_views=metrics["total_views"],
        overview=metrics["overview"],
        staff_distribution=metrics["staff_distribution"],
        at_risk_learners=metrics["at_risk_learners"],
        interventions=metrics["interventions"],
        generation_date=generation_date,
        generation_timestamp=generation_timestamp,
        faculty_name=g.user["name"],
        department=dept_name,
    )
    response = Response(html_content, mimetype="text/html; charset=utf-8")
    response.headers["Content-Disposition"] = f'attachment; filename="omnistudy_executive_report_{dept_slug}_{generation_timestamp}.html"'
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


@app.route("/documents")
@login_required()
def documents():
    query = request.args.get("q", "").strip()
    return render_template("documents.html", documents=accessible_documents(query), query=query)


@app.route("/documents/new", methods=["GET", "POST"])
@login_required()
def new_document():
    if request.method == "POST":
        access_mode = request.form.get("access_mode", "all")
        requested_students = request.form.getlist("student_ids")
        try:
            if access_mode == "block_list" and request.form.get("visibility") != "public":
                raise ValueError("Blocked-student access is available only for public repository materials.")
            validate_document_access(access_mode, requested_students)
        except ValueError as error:
            flash(str(error), "error")
            return render_template("document_form.html", document=None, students=available_students(), selected_student_ids=[])
        try:
            uploaded_content, source_filename, stored_filename = extract_upload(request.files.get("study_file"))
        except ValueError as error:
            flash(str(error), "error")
            return render_template("document_form.html", document=None, students=available_students(), selected_student_ids=[])
        content = request.form["content"].strip() or uploaded_content
        title = request.form["title"].strip()
        if not title or not content:
            flash("Provide a title and either paste study content or upload a readable file.", "error")
        else:
            db = get_db()
            cursor = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, parent_id, created_at, source_filename, stored_filename) VALUES (?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?, ?)",
                (title, request.form["subject"].strip(), request.form["department"].strip(), request.form["semester"].strip(), content, request.form["visibility"], int("release" in request.form), g.user["id"], now_iso(), source_filename, stored_filename),
            )
            db.commit()
            save_document_access(cursor.lastrowid, access_mode, requested_students)
            queue_document_processing(cursor.lastrowid)
            log_activity("document_publish", document_id=cursor.lastrowid, detail=title)
            flash("Your study material is now in OmniStudy.", "success")
            return redirect(url_for("documents"))
    return render_template("document_form.html", document=None, students=available_students(), selected_student_ids=[])


@app.route("/documents/<int:document_id>")
@login_required()
def document_detail(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    versions = [version for version in get_db().execute(
        "SELECT d.*, u.name AS owner_name FROM documents d JOIN users u ON u.id = d.owner_id WHERE d.parent_id = ? ORDER BY d.created_at DESC", (document_id,)
    ).fetchall() if visible_document(version)]
    log_activity("document_view", document_id=document_id, detail=document["title"])
    summary = stored_summary(document_id)
    content_str = document["content"] or ""
    word_count = len(content_str.split())
    read_time = max(1, round(word_count / 180))
    raw_words = [w.title() for w in re.findall(r"[A-Za-z]{4,}", content_str) if w.lower() not in STOP_WORDS and w.lower() not in STOP_CONCEPTS]
    key_terms = [item[0] for item in Counter(raw_words).most_common(8)]
    return render_template(
        "document_detail.html",
        document=document,
        summary=summary,
        versions=versions,
        shared_question_count=len(stored_questions(document_id)),
        word_count=word_count,
        read_time=read_time,
        key_terms=key_terms,
    )


@app.post("/documents/<int:document_id>/study-pack/generate")
@login_required()
def generate_shared_study_pack(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    if document["processing_status"] in {"queued", "processing"}:
        flash("The shared AI study pack is already being prepared.", "success")
    elif stored_summary(document_id) and len(stored_questions(document_id)) > 0:
        flash("This note already has a shared AI study pack. Everyone sees the saved version.", "success")
    else:
        process_document(document_id)
        flash("The shared AI summary and assessment questions are ready.", "success")
    target = request.referrer or url_for("quiz", document_id=document_id)
    return redirect(target)


@app.route("/documents/<int:document_id>/download")
@login_required()
def download_document(document_id):
    document = get_document(document_id)
    if not visible_document(document) or not document["stored_filename"]:
        abort(404)
    log_activity("document_download", document_id=document_id, detail=document["source_filename"])
    return send_from_directory(app.config["UPLOAD_FOLDER"], document["stored_filename"], as_attachment=True, download_name=document["source_filename"])


@app.route("/documents/<int:document_id>/branch", methods=["GET", "POST"])
@login_required()
def branch_document(document_id):
    original = get_document(document_id)
    if not visible_document(original):
        abort(403)
    if request.method == "POST":
        access_mode = request.form.get("access_mode", "all")
        requested_students = request.form.getlist("student_ids")
        try:
            if access_mode == "block_list" and request.form.get("visibility") != "public":
                raise ValueError("Blocked-student access is available only for public repository materials.")
            validate_document_access(access_mode, requested_students)
        except ValueError as error:
            flash(str(error), "error")
            return render_template("document_form.html", document=original, students=available_students(), selected_student_ids=[])
        try:
            uploaded_content, source_filename, stored_filename = extract_upload(request.files.get("study_file"))
        except ValueError as error:
            flash(str(error), "error")
            return render_template("document_form.html", document=original, students=available_students(), selected_student_ids=[])
        content = request.form["content"].strip() or uploaded_content
        if not request.form["title"].strip() or not content:
            flash("Provide a title and either paste revised content or upload a readable file.", "error")
        else:
            db = get_db()
            cursor = db.execute(
                "INSERT INTO documents (title, subject, department, semester, content, visibility, release_on_inactivity, owner_id, parent_id, created_at, source_filename, stored_filename) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (request.form["title"].strip(), original["subject"], original["department"], original["semester"], content, request.form["visibility"], int("release" in request.form), g.user["id"], original["id"], now_iso(), source_filename, stored_filename),
            )
            db.commit()
            save_document_access(cursor.lastrowid, access_mode, requested_students)
            queue_document_processing(cursor.lastrowid)
            log_activity("document_branch", document_id=original["id"], detail=request.form["title"].strip())
            flash("Improved version created. Compare it from the original document.", "success")
            return redirect(url_for("document_detail", document_id=document_id))
    selected_ids = selected_student_ids(original["id"], "allow" if original["access_mode"] == "allow_list" else "block")
    return render_template("document_form.html", document=original, students=available_students(), selected_student_ids=selected_ids)


@app.route("/documents/<int:document_id>/quiz", methods=["GET", "POST"])
@login_required()
def quiz(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    question_bank = stored_questions(document_id)
    full_bank_ready = len(question_bank) > 0 and document["processing_status"] not in {"queued", "processing"}
    focus = request.values.get("difficulty", "balanced")
    if focus not in QUIZ_FOCUSES:
        focus = "balanced"
    
    bank_size = len(question_bank)
    if bank_size <= 5:
        available_counts = sorted(list({c for c in [2, 3, 5, bank_size] if 1 <= c <= bank_size}))
    elif bank_size <= 12:
        available_counts = sorted(list({c for c in [3, 5, 8, 10, bank_size] if 1 <= c <= bank_size}))
    elif bank_size <= 20:
        available_counts = sorted(list({c for c in [5, 10, 15, bank_size] if 1 <= c <= bank_size}))
    else:
        available_counts = sorted(list({c for c in [5, 10, 15, 20, 25, 30] if c <= bank_size} | {bank_size}))
    if not available_counts and question_bank:
        available_counts = [bank_size]

    raw_count = request.values.get("question_count")
    requested_count = quiz_count(raw_count, bank_size)
    if requested_count not in available_counts and available_counts:
        requested_count = min(available_counts, key=lambda c: abs(c - requested_count))

    seed_val = request.values.get("seed")
    if not seed_val:
        seed_val = str(random.randint(100000, 999999))
    next_seed = str(random.randint(100000, 999999))

    questions = select_quiz_questions(question_bank, focus, requested_count, seed=seed_val) if full_bank_ready and requested_count else []
    result = None
    if request.method == "POST" and questions:
        submitted_answers = [request.form.get(f"question_{index}", "").strip() for index in range(len(questions))]
        score = sum(submitted_answers[index] == question["answer"].strip() for index, question in enumerate(questions))
        get_db().execute(
            "INSERT INTO quiz_attempts (document_id, user_id, score, total, created_at) VALUES (?, ?, ?, ?, ?)",
            (document_id, g.user["id"], score, len(questions), now_iso()),
        )
        get_db().commit()
        log_activity("quiz_complete", document_id=document_id, detail=f"{score}/{len(questions)}")
        mastery = round(score * 100 / len(questions))
        message = "Excellent retrieval practice." if mastery >= 80 else "A useful baseline—review the explanations and try again." if mastery < 50 else "You are building strong momentum."
        result = {"score": score, "total": len(questions), "mastery": mastery, "message": message, "answers": submitted_answers}
    
    db = get_db()
    student_document_attempts = db.execute(
        "SELECT COUNT(*) FROM quiz_attempts WHERE document_id = ? AND user_id = ?",
        (document_id, g.user["id"])
    ).fetchone()[0]
    return render_template(
        "quiz.html", document=document, questions=questions, result=result, difficulty=focus,
        requested_count=requested_count, available_counts=available_counts, bank_size=bank_size,
        full_bank_ready=full_bank_ready, overall_performance=learner_performance(g.user["id"]),
        student_document_attempts=student_document_attempts, seed=seed_val, next_seed=next_seed,
    )



@app.post("/documents/<int:document_id>/quiz/attempts/clear")
@login_required()
def clear_document_quiz_attempts(document_id):
    db = get_db()
    db.execute("DELETE FROM quiz_attempts WHERE document_id = ? AND user_id = ?", (document_id, g.user["id"]))
    db.commit()
    flash("Your previous quiz attempts on this note have been cleared.", "success")
    return redirect(url_for("quiz", document_id=document_id))


@app.post("/documents/<int:document_id>/study-pack/refresh")
@login_required()
def refresh_study_pack(document_id):
    document = get_document(document_id)
    if document["owner_id"] != g.user["id"] and g.user["role"] != "admin":
        abort(403)
    queue_document_processing(document_id)
    flash("Your AI study pack is being refreshed.", "success")
    return redirect(url_for("document_detail", document_id=document_id))


@app.post("/documents/<int:document_id>/quiz/delete")
@login_required()
def delete_study_pack(document_id):
    document = get_document(document_id)
    if not document:
        abort(404)
    if document["owner_id"] != g.user["id"] and g.user["role"] != "admin":
        abort(403)
    db = get_db()
    db.execute("DELETE FROM quiz_questions WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM quiz_attempts WHERE document_id = ?", (document_id,))
    db.commit()
    log_activity("quiz_delete", document_id=document_id, detail=f"Deleted question bank for {document['title']}")
    flash("The generated quiz question bank has been deleted. You can generate a fresh one anytime.", "success")
    target = request.referrer or url_for("quiz", document_id=document_id)
    return redirect(target)


@app.post("/documents/<int:document_id>/delete")
@login_required()
def delete_document(document_id):
    document = get_document(document_id)
    if not document:
        abort(404)
    if document["owner_id"] != g.user["id"] and g.user["role"] != "admin":
        abort(403)
    db = get_db()
    if document["stored_filename"]:
        file_path = os.path.join(app.config["UPLOAD_FOLDER"], document["stored_filename"])
        if os.path.exists(file_path):
            try:
                os.remove(file_path)
            except Exception:
                pass
    db.execute("DELETE FROM quiz_questions WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM study_summaries WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM quiz_attempts WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM document_access WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM flashcard_progress WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM viva_history WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM activity_logs WHERE document_id = ?", (document_id,))
    db.execute("UPDATE documents SET parent_id = NULL WHERE parent_id = ?", (document_id,))
    doc_title = document["title"]
    db.execute("DELETE FROM documents WHERE id = ?", (document_id,))
    db.commit()
    log_activity("document_delete", document_id=None, detail=f"Deleted material: {doc_title}")
    flash(f"Material \"{document['title']}\" was deleted successfully.", "success")
    return redirect(url_for("documents"))


@app.post("/quiz/attempts/<int:attempt_id>/delete")
@login_required()
def delete_quiz_attempt(attempt_id):
    db = get_db()
    attempt = db.execute("SELECT * FROM quiz_attempts WHERE id = ?", (attempt_id,)).fetchone()
    if not attempt:
        abort(404)
    if attempt["user_id"] != g.user["id"] and g.user["role"] != "admin":
        abort(403)
    db.execute("DELETE FROM quiz_attempts WHERE id = ?", (attempt_id,))
    db.commit()
    flash("Quiz attempt removed from history.", "success")
    return redirect(request.referrer or url_for("dashboard"))


@app.route("/documents/<int:document_id>/engagement")
@login_required()
def document_engagement(document_id):
    document = get_document(document_id)
    if not document or document["owner_id"] != g.user["id"]:
        abort(403)
    db = get_db()
    viewer_events = db.execute(
        "SELECT u.name, u.email, u.department, u.semester, a.created_at "
        "FROM activity_logs a JOIN users u ON u.id = a.user_id "
        "WHERE a.document_id = ? AND a.event_type = 'document_view' AND u.role = 'student' AND u.id != ? "
        "ORDER BY a.created_at DESC",
        (document_id, g.user["id"]),
    ).fetchall()
    unique_students = len({event["email"] for event in viewer_events})
    return render_template(
        "document_engagement.html",
        document=document,
        viewer_events=viewer_events,
        unique_students=unique_students,
        total_views=len(viewer_events),
    )


@app.route("/documents/<int:original_id>/compare/<int:version_id>")
@login_required()
def compare_documents(original_id, version_id):
    original = get_document(original_id)
    version = get_document(version_id)
    if not visible_document(original) or not visible_document(version) or version["parent_id"] != original_id:
        abort(403)
    diff_data = comparison_rows(original["content"], version["content"])
    stats = getattr(diff_data, "stats", {})
    return render_template(
        "compare.html",
        original=original,
        version=version,
        rows=diff_data,
        stats=stats
    )


@app.route("/staff/analytics")
@login_required({"staff", "admin"})
def staff_analytics():
    db = get_db()
    metrics = get_staff_comprehensive_metrics(g.user)
    staff_doc_ids = get_staff_accessible_doc_ids(g.user)
    placeholders = ",".join("?" for _ in staff_doc_ids) if staff_doc_ids else "?"
    params = staff_doc_ids if staff_doc_ids else [-1]

    # Viva oral metrics per user
    viva_map = {}
    try:
        viva_rows = db.execute(f"""
            SELECT user_id, COUNT(*) AS cnt, AVG(score) AS avg_s
            FROM viva_history
            WHERE document_id IN ({placeholders})
            GROUP BY user_id
        """, params).fetchall()
        for vr in viva_rows:
            viva_map[vr["user_id"]] = {
                "count": vr["cnt"],
                "avg_score": round(vr["avg_s"] or 0, 1)
            }
    except Exception:
        pass

    # Flashcard reviews per user
    flashcard_map = {}
    try:
        fc_rows = db.execute(f"""
            SELECT user_id, COUNT(*) AS cnt
            FROM flashcard_progress
            WHERE document_id IN ({placeholders})
            GROUP BY user_id
        """, params).fetchall()
        for fcr in fc_rows:
            flashcard_map[fcr["user_id"]] = fcr["cnt"]
    except Exception:
        pass

    # Student per-document attempt breakdown
    breakdowns = {}
    try:
        doc_rows = db.execute(f"""
            SELECT q.user_id, d.id AS doc_id, d.title AS doc_title, d.subject,
                   COUNT(q.id) AS attempts,
                   ROUND(AVG(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END), 1) AS avg_score,
                   MAX(ROUND(CASE WHEN q.total > 0 THEN q.score * 100.0 / q.total ELSE 0 END, 1)) AS best_score
            FROM quiz_attempts q
            JOIN documents d ON d.id = q.document_id
            WHERE q.document_id IN ({placeholders})
            GROUP BY q.user_id, d.id
            ORDER BY d.title ASC
        """, params).fetchall()
        for dr in doc_rows:
            breakdowns.setdefault(dr["user_id"], []).append({
                "doc_id": dr["doc_id"],
                "title": dr["doc_title"],
                "subject": dr["subject"],
                "attempts": dr["attempts"],
                "avg_score": dr["avg_score"] or 0,
                "best_score": dr["best_score"] or 0
            })
    except Exception:
        pass

    analytics = [
        {
            "id": n["id"],
            "title": n["title"],
            "subject": n["subject"],
            "created_at": n["created_at"],
            "views": n["views_count"],
            "learners": n["unique_students_count"],
            "attempts": n["quiz_attempts_count"],
            "average_score": n["avg_mastery"],
        }
        for n in metrics["staff_notes"]
    ]
    overview = metrics["overview"]
    
    learner_results = []
    at_risk_list = []
    
    for l in metrics["staff_learners"]:
        uid = l["id"]
        avg_score = l["avg_score"] or 0
        best_score = l["max_score"] or 0
        viva = viva_map.get(uid, {"count": 0, "avg_score": 0})
        fc_count = flashcard_map.get(uid, 0)
        docs = breakdowns.get(uid, [])
        
        # Weakest topic for targeted remediation
        weakest_doc = None
        if docs:
            sorted_by_score = sorted([d for d in docs if d["avg_score"] is not None], key=lambda x: x["avg_score"])
            if sorted_by_score:
                weakest_doc = sorted_by_score[0]

        tier = "A" if avg_score >= 85 else ("B" if avg_score >= 70 else ("C" if avg_score >= 50 else "D"))
        is_at_risk = avg_score < 60
        
        item = {
            "id": uid,
            "name": l["name"],
            "email": l["email"],
            "department": l["department"] if "department" in l.keys() and l["department"] else "BCA",
            "semester": l["semester"] if "semester" in l.keys() and l["semester"] else "6",
            "attempts": l["total_attempts"],
            "average_score": avg_score,
            "best_score": best_score,
            "last_attempt": l["last_active"],
            "tier": tier,
            "is_at_risk": is_at_risk,
            "viva_count": viva["count"],
            "viva_score": viva["avg_score"],
            "flashcards_reviewed": fc_count,
            "doc_breakdown": docs,
            "weakest_topic": weakest_doc["title"] if weakest_doc else "General Revision",
        }
        learner_results.append(item)
        if is_at_risk:
            at_risk_list.append(item)

    dist = metrics["staff_distribution"]
    total_grades = max(1, dist["total"])
    dist_pct = {
        "A": round((dist["A"] / total_grades) * 100, 1),
        "B": round((dist["B"] / total_grades) * 100, 1),
        "C": round((dist["C"] / total_grades) * 100, 1),
        "D": round((dist["D"] / total_grades) * 100, 1),
        "total": dist["total"]
    }

    return render_template(
        "staff_analytics.html",
        analytics=analytics,
        overview=overview,
        learner_results=learner_results,
        at_risk_learners=at_risk_list,
        distribution=dist,
        distribution_pct=dist_pct,
        knowledge_gaps=metrics["staff_gaps"],
    )


@app.route("/admin", methods=["GET", "POST"])
@login_required({"admin"})
def admin():
    db = get_db()
    if request.method == "POST":
        try:
            threshold = max(1, int(request.form.get("inactivity_days", 0)))
        except (ValueError, TypeError):
            flash("Inactivity days must be a positive whole number.", "error")
        else:
            db.execute("UPDATE system_settings SET setting_value = ? WHERE setting_key = 'inactivity_days'", (str(threshold),))
            db.commit()
            released = run_inactivity_watchdog()
            flash(f"Settings saved. Watchdog completed: {released} deferred documents released.", "success")
    settings = dict(db.execute("SELECT setting_key, setting_value FROM system_settings"))
    users = db.execute("SELECT id, name, reg_no, email, role, last_login, is_verified, graduation_status FROM users ORDER BY is_verified, last_login DESC").fetchall()
    pending_profile_requests = db.execute(
        "SELECT p.*, u.email as user_email FROM profile_update_requests p JOIN users u ON u.id = p.user_id WHERE p.status = 'pending' ORDER BY p.created_at DESC"
    ).fetchall()
    events = db.execute(
        "SELECT a.event_type, a.detail, a.created_at, u.name FROM activity_logs a JOIN users u ON u.id = a.user_id ORDER BY a.created_at DESC LIMIT 12"
    ).fetchall()
    stored_uploads_count = str(len(os.listdir(UPLOAD_FOLDER))) if os.path.exists(UPLOAD_FOLDER) else "0"
    health = {
        "Database": "Connected" if os.path.exists(DATABASE) else "Unavailable",
        "Background watchdog": "Running" if scheduler.running else "Starts with server",
        "Stored uploads": stored_uploads_count,
    }
    return render_template("admin.html", settings=settings, users=users, pending_profile_requests=pending_profile_requests, events=events, health=health)


@app.post("/admin/users/<int:user_id>/verify")
@login_required({"admin"})
def verify_user(user_id):
    db = get_db()
    user = db.execute("SELECT id, role, is_verified FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user or user["role"] == "admin":
        abort(404)
    if not user["is_verified"]:
        db.execute("UPDATE users SET is_verified = 1 WHERE id = ?", (user_id,))
        db.execute("INSERT INTO verification_logs (user_id, admin_id, action, created_at) VALUES (?, ?, 'verified', ?)", (user_id, g.user["id"], now_iso()))
        db.execute("INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'verification', 'Account approved by administrator', ?)", (user_id, now_iso()))
        db.commit()
        flash("Account verified.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/users/<int:user_id>/graduation")
@login_required({"admin"})
def update_graduation_status(user_id):
    db = get_db()
    user = db.execute("SELECT id, role FROM users WHERE id = ?", (user_id,)).fetchone()
    status = request.form.get("graduation_status")
    if not user or user["role"] != "student" or status not in {"active", "graduated"}:
        abort(404)
    db.execute("UPDATE users SET graduation_status = ? WHERE id = ?", (status, user_id))
    db.execute("INSERT INTO verification_logs (user_id, admin_id, action, created_at) VALUES (?, ?, ?, ?)", (user_id, g.user["id"], f"graduation_status:{status}", now_iso()))
    db.commit()
    flash("Student lifecycle status updated.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/profile-requests/<int:request_id>/approve")
@login_required({"admin"})
def approve_profile_request(request_id):
    db = get_db()
    req = db.execute("SELECT * FROM profile_update_requests WHERE id = ?", (request_id,)).fetchone()
    if not req or req["status"] != "pending":
        abort(404)
    db.execute(
        "UPDATE users SET name = ?, reg_no = ?, department = ?, semester = ? WHERE id = ?",
        (req["new_name"], req["new_reg_no"], req["new_department"], req["new_semester"], req["user_id"])
    )
    db.execute(
        "UPDATE profile_update_requests SET status = 'approved', admin_id = ?, reviewed_at = ? WHERE id = ?",
        (g.user["id"], now_iso(), request_id)
    )
    db.execute(
        "INSERT INTO verification_logs (user_id, admin_id, action, created_at) VALUES (?, ?, 'profile_update_approved', ?)",
        (req["user_id"], g.user["id"], now_iso())
    )
    db.execute(
        "INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'profile_update', 'Profile changes approved by administrator', ?)",
        (req["user_id"], now_iso())
    )
    db.commit()
    flash(f"Profile update approved for {req['new_name']}.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/profile-requests/<int:request_id>/reject")
@login_required({"admin"})
def reject_profile_request(request_id):
    db = get_db()
    req = db.execute("SELECT * FROM profile_update_requests WHERE id = ?", (request_id,)).fetchone()
    if not req or req["status"] != "pending":
        abort(404)
    db.execute(
        "UPDATE profile_update_requests SET status = 'rejected', admin_id = ?, reviewed_at = ? WHERE id = ?",
        (g.user["id"], now_iso(), request_id)
    )
    db.execute(
        "INSERT INTO verification_logs (user_id, admin_id, action, created_at) VALUES (?, ?, 'profile_update_rejected', ?)",
        (req["user_id"], g.user["id"], now_iso())
    )
    db.execute(
        "INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'profile_update', 'Profile update request rejected by administrator', ?)",
        (req["user_id"], now_iso())
    )
    db.commit()
    flash("Profile update request rejected.", "success")
    return redirect(url_for("admin"))


@app.route("/profile")
@login_required()
def profile():
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id = ?", (g.user["id"],)).fetchone()
    pending_request = db.execute(
        "SELECT * FROM profile_update_requests WHERE user_id = ? AND status = 'pending' ORDER BY created_at DESC LIMIT 1",
        (g.user["id"],)
    ).fetchone()
    past_requests = db.execute(
        "SELECT * FROM profile_update_requests WHERE user_id = ? AND status != 'pending' ORDER BY created_at DESC LIMIT 5",
        (g.user["id"],)
    ).fetchall()
    return render_template("profile.html", user=user, pending_request=pending_request, past_requests=past_requests)


@app.post("/profile/edit")
@login_required()
def edit_profile():
    new_name = request.form.get("name", "").strip()
    new_reg_no = request.form.get("reg_no", "").strip()
    new_dept = request.form.get("department", "").strip()
    new_sem = request.form.get("semester", "").strip()

    if not new_name or not new_dept or not new_sem:
        flash("Name, department, and semester are required.", "error")
        return redirect(url_for("profile"))

    db = get_db()
    current = db.execute("SELECT * FROM users WHERE id = ?", (g.user["id"],)).fetchone()
    current_reg = current["reg_no"] or ""

    if (new_name == current["name"] and 
        new_reg_no == current_reg and 
        new_dept == current["department"] and 
        new_sem == current["semester"]):
        flash("No changes detected. Your details are already up to date.", "error")
        return redirect(url_for("profile"))

    existing = db.execute(
        "SELECT id FROM profile_update_requests WHERE user_id = ? AND status = 'pending'",
        (g.user["id"],)
    ).fetchone()

    if existing:
        db.execute(
            """UPDATE profile_update_requests 
               SET new_name = ?, new_reg_no = ?, new_department = ?, new_semester = ?, created_at = ?
               WHERE id = ?""",
            (new_name, new_reg_no, new_dept, new_sem, now_iso(), existing["id"])
        )
        flash("Your existing pending profile update request has been updated with the new details.", "success")
    else:
        db.execute(
            """INSERT INTO profile_update_requests 
               (user_id, current_name, new_name, current_reg_no, new_reg_no, current_department, new_department, current_semester, new_semester, status, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?)""",
            (g.user["id"], current["name"], new_name, current_reg, new_reg_no, current["department"], new_dept, current["semester"], new_sem, now_iso())
        )
        db.execute(
            "INSERT INTO activity_logs (user_id, event_type, detail, created_at) VALUES (?, 'profile_request', 'Submitted profile edit request for admin approval', ?)",
            (g.user["id"], now_iso())
        )
        flash("Profile update request submitted successfully. It will take effect once approved by an administrator.", "success")

    db.commit()
    return redirect(url_for("profile"))


@app.post("/profile/cancel-request")
@login_required()
def cancel_profile_request():
    db = get_db()
    existing = db.execute(
        "SELECT id FROM profile_update_requests WHERE user_id = ? AND status = 'pending'",
        (g.user["id"],)
    ).fetchone()
    if existing:
        db.execute("DELETE FROM profile_update_requests WHERE id = ?", (existing["id"],))
        db.commit()
        flash("Your pending profile update request has been cancelled.", "success")
    return redirect(url_for("profile"))


# ==========================================
# ADVANCED INNOVATION SUITE (4 NEW ENGINES)
# ==========================================

def build_flashcards(content):
    lines = content.splitlines()
    cards = []
    seen_fronts = set()

    def add_card(front, back, category, hint=""):
        f_clean = front.strip()
        b_clean = back.strip()
        if f_clean.lower() in seen_fronts or len(f_clean) < 3 or len(b_clean) < 5:
            return
        seen_fronts.add(f_clean.lower())
        cid = hashlib.md5(f_clean.encode('utf-8')).hexdigest()[:10]
        cards.append({
            "id": cid,
            "front": f_clean,
            "back": b_clean,
            "category": category,
            "hint": hint or (b_clean[:45] + "..." if len(b_clean) > 45 else b_clean)
        })

    # 1. Structure extraction (Topic -> Role / Warning)
    current_topic = None
    for line in lines:
        header_m = re.match(r"^([A-Z0-9\s&/\-\(\)]{3,40})$", line)
        if header_m and not line.lower().startswith(("role", "habit", "watch", "page", "tip", "note", "source", "below", "daily")):
            current_topic = header_m.group(1).strip().title()
            continue
        role_m = re.match(r"^Role(?:\s+in\s+the\s+body)?:\s*(.+)$", line, re.IGNORECASE)
        if role_m and current_topic:
            add_card(f"Primary Function of {current_topic}", role_m.group(1).strip(), "Core Function")
        watch_m = re.match(r"^Watch\s+out\s+for:\s*(.+)$", line, re.IGNORECASE)
        if watch_m and current_topic:
            add_card(f"Clinical Warning Signs: {current_topic}", watch_m.group(1).strip(), "Warning Signs")

    # 2. Key-Value definitions
    for line in lines:
        kv_m = re.match(r"^([A-Z][a-zA-Z0-9\s\(\)\-\.,]{2,30}?)\s+(?:—|–|:)\s+(.+)$", line)
        if kv_m:
            term = kv_m.group(1).strip()
            desc = kv_m.group(2).strip()
            if not any(term.lower().startswith(sc) for sc in STOP_CONCEPTS):
                add_card(f"Define: {term}", desc, "Definition")

    # 3. Nutrients table
    nutrient_names = ["Vitamin A", "Vitamin B1 (Thiamine)", "Vitamin B2 (Riboflavin)", "Vitamin B3 (Niacin)", "Vitamin B6", "Vitamin B9 (Folate)", "Vitamin B12", "Vitamin C", "Vitamin D", "Vitamin E", "Vitamin K", "Calcium", "Iron", "Magnesium", "Zinc", "Potassium"]
    for idx, l in enumerate(lines):
        for n in nutrient_names:
            if l.lower().startswith(n.lower()):
                rem = l[len(n):].strip().lstrip("—–:- ").strip()
                if len(rem) > 10:
                    add_card(f"Health Role of {n}", rem, "Biochemistry")
                elif idx + 1 < len(lines) and len(lines[idx+1].strip()) > 10:
                    add_card(f"Health Role of {n}", lines[idx+1].strip(), "Biochemistry")
                break

    # 4. Academic verbs sentence extraction
    academic_verbs = r"is\s+(?:the|a|an)?|are|provides?|manages?|enables?|regulates?|supports?|solves?|stores?|coordinates?|optimizes?|processes?"
    for line in lines:
        sm = re.match(rf"^(?:The\s+|A\s+)?([A-Z][a-zA-Z0-9_\-\s]{{2,30}}?)\s+({academic_verbs})\s+(.+)$", line, re.IGNORECASE)
        if sm:
            term = sm.group(1).strip().title()
            verb = sm.group(2).strip()
            pred = sm.group(3).strip()
            if term.lower() not in STOP_CONCEPTS and len(term.split()) <= 4 and len(pred) > 10:
                add_card(f"What is the mechanism of {term}?", f"{verb.capitalize()} {pred}", "Mechanism")

    if not cards:
        sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", content) if len(s.strip()) > 30]
        for idx, s in enumerate(sentences[:10]):
            add_card(f"Key Concept {idx+1}", s, "Study Point")

    return cards


def extract_concept_graph(content, doc_title="Subject Domain"):
    lines = [l.strip() for l in content.splitlines() if l.strip()]
    nodes = {}
    edges = []
    seen_edges = set()
    
    stop_concepts = {
        "the", "and", "or", "an", "a", "in", "on", "at", "to", "for", "with", "from",
        "by", "of", "as", "is", "are", "was", "were", "be", "been", "that", "this",
        "these", "those", "it", "its", "they", "their", "which", "what", "how", "all",
        "each", "every", "some", "any", "not", "also", "into", "over", "such", "than",
        "note", "notes", "chapter", "section", "lecture", "overview", "summary", "introduction",
        "guide", "complete", "daily", "below", "following", "steps", "points", "drawbacks", "purpose",
        "draw this", "diagram", "table", "exam tip", "source", "key"
    }

    def add_node(name, group="entity", desc="", icon="📌", parent=None, quote=""):
        clean = re.sub(r'^[#*\-•\d\.\s]+', '', name).strip()
        clean = re.sub(r'\s*\(.*?\)', '', clean).strip()
        if len(clean) < 3 or clean.lower() in stop_concepts or len(clean.split()) > 4:
            return None
        nid = re.sub(r'[^a-zA-Z0-9]', '_', clean).lower()
        if nid not in nodes:
            base_val = 34 if group == 'core' else (25 if group == 'cluster' else (18 if group == 'entity' else 14))
            nodes[nid] = {
                "id": nid,
                "label": clean.title(),
                "group": group,  # 'core', 'cluster', 'entity', 'property'
                "desc": desc or f"Key concept in {doc_title}",
                "quote": quote or desc[:140] if desc else f"Document reference for {clean.title()}",
                "icon": icon,
                "val": base_val,
                "connections": 0,
                "importance": 1.0,
            }
        elif desc and len(desc) > len(nodes[nid]["desc"]):
            nodes[nid]["desc"] = desc
            if quote:
                nodes[nid]["quote"] = quote
        if parent:
            add_edge(parent, nid, "contains", "#6366f1", "structural")
        return nid

    def add_edge(source, target, label="relates to", color="#6366f1", edge_type="structural"):
        if not source or not target or source == target:
            return
        if source not in nodes or target not in nodes:
            return
        edge_key = f"{source}->{target}:{label}"
        if edge_key in seen_edges:
            return
        seen_edges.add(edge_key)
        nodes[source]["connections"] += 1
        nodes[target]["connections"] += 1
        edges.append({
            "source": source,
            "target": target,
            "label": label,
            "color": color,
            "type": edge_type
        })

    # 1. Core Domain Root
    core_id = add_node(doc_title, "core", f"Central Subject Domain: {doc_title}", "👑", quote=f"Central subject domain: {doc_title}")
    
    # 2. Structure & Definition Parsing
    current_cluster = core_id
    for line in lines:
        # Module / Section Headers
        header_m = re.match(r"^(?:#{1,3}\s+)?([A-Z0-9\s&/\-\(\)]{3,35})$", line)
        if header_m and not line.lower().startswith(("role", "habit", "watch", "tip", "note", "source", "below", "daily", "definition", "guide")):
            sec_title = header_m.group(1).strip()
            if len(sec_title.split()) <= 4 and sec_title.lower() != doc_title.lower():
                cluster_id = add_node(sec_title, "cluster", f"Topic Area: {sec_title}", "📦", core_id, quote=line)
                if cluster_id:
                    current_cluster = cluster_id
                    continue

        # Definitions: "Term — Definition" or "Term : Definition"
        def_m = re.match(r"^([A-Z][a-zA-Z0-9\s\(\)\-\.,]{2,30}?)\s+(?:—|–|:)\s+(.+)$", line)
        if def_m:
            term = def_m.group(1).strip()
            desc = def_m.group(2).strip()
            tid = add_node(term, "entity", desc, "📌", current_cluster or core_id, quote=line)
            if tid:
                keywords = re.findall(r"\b([A-Z][a-zA-Z0-9]{3,})\b", desc)
                for kw in keywords[:2]:
                    if kw.lower() not in stop_concepts and kw.lower() != term.lower():
                        kid = add_node(kw, "property", f"Key aspect of {term}", "⚙️", quote=desc)
                        if kid:
                            add_edge(tid, kid, "specifies", "#06b6d4", "definitional")

        # Academic Verb Relational Sentences
        academic_verbs_pattern = (
            r"is|are|provides|manages|enables|supports|regulates|solves|prevents|"
            r"stores|optimizes|consists of|contains|implements|requires|produces|transforms|mitigates"
        )
        rel_pattern = rf"^(?:The\s+|A\s+)?([A-Z][a-zA-Z0-9_\-\s]{{2,28}}?)\s+({academic_verbs_pattern})\s+(.+)$"
        sm = re.match(rel_pattern, line, re.IGNORECASE)
        if sm:
            src = sm.group(1).strip()
            rel = sm.group(2).strip().lower()
            tgt_text = sm.group(3).strip()
            src_id = add_node(src, "entity", f"{src} {rel} {tgt_text[:60]}...", "⚙️", current_cluster or core_id, quote=line)
            
            tgt_match = re.findall(r"\b([A-Z][a-zA-Z0-9]{3,})\b", tgt_text)
            for tm in tgt_match[:2]:
                if tm.lower() not in stop_concepts and tm.lower() != src.lower():
                    tgt_id = add_node(tm, "property", f"{tm} in relation to {src}", "🔹", quote=line)
                    if src_id and tgt_id:
                        if rel in ("solves", "prevents", "mitigates"):
                            edge_color = "#f59e0b"
                            edge_type = "causal"
                        elif rel in ("requires", "depends on"):
                            edge_color = "#ec4899"
                            edge_type = "dependency"
                        elif rel in ("provides", "enables", "optimizes", "implements", "produces"):
                            edge_color = "#10b981"
                            edge_type = "functional"
                        else:
                            edge_color = "#8b5cf6"
                            edge_type = "relational"
                        add_edge(src_id, tgt_id, rel, edge_color, edge_type)

    # Connect any disconnected nodes to core hub
    for nid, node in list(nodes.items()):
        if nid != core_id and node["connections"] == 0:
            add_edge(core_id, nid, "covers", "#6366f1", "structural")

    # Dynamic importance and radius sizing
    for nid, node in nodes.items():
        conn = node["connections"]
        if node["group"] == "core":
            node["val"] = 34
            node["importance"] = 3.0
        elif node["group"] == "cluster":
            node["val"] = min(30, 22 + conn * 2)
            node["importance"] = round(1.5 + conn * 0.2, 2)
        else:
            node["val"] = min(26, 14 + conn * 2)
            node["importance"] = round(1.0 + conn * 0.15, 2)

    sorted_nodes = sorted(nodes.values(), key=lambda n: (n["group"] == 'core', n["group"] == 'cluster', n["connections"]), reverse=True)[:35]
    valid_node_ids = {n["id"] for n in sorted_nodes}
    valid_edges = [e for e in edges if e["source"] in valid_node_ids and e["target"] in valid_node_ids][:50]

    return {
        "nodes": sorted_nodes,
        "edges": valid_edges,
        "metrics": {
            "total_nodes": len(sorted_nodes),
            "total_edges": len(valid_edges),
            "clusters": sum(1 for n in sorted_nodes if n["group"] == "cluster"),
            "entities": sum(1 for n in sorted_nodes if n["group"] == "entity"),
            "properties": sum(1 for n in sorted_nodes if n["group"] == "property")
        }
    }


def clean_cheat_sheet_term(text):
    t = re.sub(r"^(Define[:\s]*|What is (the mechanism of )?|Explain[:\s]*|Primary Function of\s*)", "", text, flags=re.IGNORECASE).strip()
    return t.rstrip(":?-— ") or text


def generate_cheat_sheet_data(content):
    cards = build_flashcards(content)
    
    # 1. Core definitions
    raw_defs = [c for c in cards if "Define:" in c.get("front", "")]
    if len(raw_defs) < 4:
        raw_defs = cards[:6]
    else:
        raw_defs = raw_defs[:6]

    definitions = []
    for c in raw_defs:
        term = clean_cheat_sheet_term(c.get("front", ""))
        definitions.append({
            "term": term,
            "def": c.get("back", ""),
            "weightage": "2 Marks"
        })

    # 2. Key takeaways / mechanisms
    mechanisms = [c["back"] for c in cards if c.get("category") in ("Core Function", "Mechanism", "Warning Signs")][:6]
    if len(mechanisms) < 3:
        mechanisms = [c["back"] for c in cards[len(definitions):len(definitions)+5]]
    if not mechanisms and cards:
        mechanisms = [c["back"] for c in cards[:4]]

    # 3. High-Frequency Exam Q&A
    exam_qa = []
    for c in cards[:5]:
        front = c.get("front", "")
        clean_front = clean_cheat_sheet_term(front)
        if clean_front and not clean_front.endswith("?"):
            if any(clean_front.lower().startswith(w) for w in ["what", "how", "why", "explain", "describe", "differentiate"]):
                q_text = f"{clean_front}?"
            else:
                q_text = f"What is {clean_front}?"
        else:
            q_text = clean_front or front
            
        exam_qa.append({
            "q": q_text,
            "a": c.get("back", ""),
            "weightage": "5 Marks" if len(c.get("back", "").split()) > 16 else "2 Marks"
        })

    # 4. Dynamic exam readiness checklist
    checklist = []
    for d in definitions[:4]:
        if d.get("term"):
            checklist.append(f"Can define '{d['term']}' without referring to notes")
    if mechanisms:
        checklist.append("Can summarize core operational workflow & principles")
    if not checklist:
        checklist = [
            "Can define primary core concepts without notes",
            "Can list the primary operational workflow",
            "Solved 10+ MCQ practice questions"
        ]

    return {
        "definitions": definitions,
        "mechanisms": mechanisms,
        "exam_qa": exam_qa,
        "checklist": checklist
    }


def generate_viva_questions(content):
    cards = build_flashcards(content)
    questions = []
    for idx, c in enumerate(cards[:8]):
        topic_name = c['front'].replace('Define: ', '').replace('Primary Function of ', '').replace('What is the mechanism of ', '')
        q_text = f"Can you explain '{topic_name}' in detail and describe its core operational or physiological principles?"
        keywords = [w.lower() for w in re.findall(r"[A-Za-z]{4,}", c["back"]) if w.lower() not in STOP_WORDS]
        questions.append({
            "id": idx + 1,
            "topic": topic_name,
            "question": q_text,
            "model_answer": c["back"],
            "expected_keywords": keywords[:6]
        })
    return questions


def evaluate_viva_answer(expected_keywords, model_answer, student_answer):
    s_clean = student_answer.strip().lower()
    if not s_clean or len(s_clean.split()) < 3:
        return {
            "score": 1.0,
            "grade": "Needs Review (D)",
            "matched_keywords": [],
            "missed_keywords": expected_keywords,
            "feedback": "Your response was too brief. Elaborate on the core mechanisms and include formal technical terminology."
        }
        
    student_words = set(re.findall(r"[a-z]{3,}", s_clean))
    matched = [k for k in expected_keywords if k in student_words]
    missed = [k for k in expected_keywords if k not in student_words]
    
    coverage = len(matched) / max(1, len(expected_keywords))
    length_bonus = min(1.0, len(s_clean.split()) / 25.0)
    
    raw_score = (coverage * 7.5) + (length_bonus * 2.5)
    score = round(min(10.0, max(2.0, raw_score)), 1)
    
    if score >= 8.5:
        grade = "Mastery (A+)"
        feedback = "Outstanding articulation! You demonstrated clear mastery and utilized essential domain terminology."
    elif score >= 7.0:
        grade = "Proficient (B+)"
        feedback = "Solid explanation! You covered the primary concepts with good understanding."
    elif score >= 5.0:
        grade = "Developing (C)"
        feedback = f"Good attempt, but missing key concepts. Incorporate: {', '.join(missed[:3])}."
    else:
        grade = "Needs Review (D)"
        feedback = f"Basic attempt. To elevate your score, explain the mechanisms and mention: {', '.join(missed[:3])}."

    return {
        "score": score,
        "grade": grade,
        "matched_keywords": matched,
        "missed_keywords": missed,
        "feedback": feedback
    }


# 1. 3D FLASHCARDS ROUTE
@app.route("/documents/<int:document_id>/flashcards")
@login_required()
def flashcards_view(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    cards = build_flashcards(document["content"] or "")
    db = get_db()
    progress_rows = db.execute(
        "SELECT card_key, box_level, review_count FROM flashcard_progress WHERE user_id = ? AND document_id = ?",
        (g.user["id"], document_id)
    ).fetchall()
    progress_map = {r["card_key"]: {"box": r["box_level"], "count": r["review_count"]} for r in progress_rows}
    
    enriched_cards = []
    box_counts = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}
    for c in cards:
        p = progress_map.get(c["id"], {"box": 1, "count": 0})
        box = p["box"]
        box_counts[box] = box_counts.get(box, 0) + 1
        enriched_cards.append({**c, "box": box, "review_count": p["count"]})
        
    mastery_pct = round((sum(c["box"] - 1 for c in enriched_cards) / max(1, len(enriched_cards) * 4)) * 100) if enriched_cards else 0
    return render_template(
        "flashcards.html",
        document=document,
        cards=enriched_cards,
        box_counts=box_counts,
        mastery_pct=mastery_pct,
    )


@app.post("/documents/<int:document_id>/flashcards/review")
@login_required()
def flashcard_review(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    card_id = request.form.get("card_id")
    rating = request.form.get("rating", "good")
    db = get_db()
    current = db.execute(
        "SELECT box_level, review_count FROM flashcard_progress WHERE user_id = ? AND document_id = ? AND card_key = ?",
        (g.user["id"], document_id, card_id)
    ).fetchone()
    
    cur_box = current["box_level"] if current else 1
    cur_count = current["review_count"] if current else 0
    
    if rating == "again":
        new_box = 1
    elif rating == "hard":
        new_box = max(1, cur_box)
    elif rating == "good":
        new_box = min(5, cur_box + 1)
    elif rating == "easy":
        new_box = min(5, cur_box + 2)
    else:
        new_box = cur_box
        
    db.execute(
        "INSERT INTO flashcard_progress (user_id, document_id, card_key, box_level, review_count, last_reviewed) "
        "VALUES (?, ?, ?, ?, ?, ?) "
        "ON CONFLICT(user_id, document_id, card_key) DO UPDATE SET "
        "box_level = excluded.box_level, review_count = review_count + 1, last_reviewed = excluded.last_reviewed",
        (g.user["id"], document_id, card_id, new_box, cur_count + 1, now_iso())
    )
    db.commit()
    log_activity("flashcard_review", document_id=document_id, detail=f"Card {card_id} -> Box {new_box}")
    return jsonify({"success": True, "card_id": card_id, "new_box": new_box})


@app.post("/documents/<int:document_id>/flashcards/reset")
@login_required()
def flashcards_reset(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    db = get_db()
    db.execute("DELETE FROM flashcard_progress WHERE user_id = ? AND document_id = ?", (g.user["id"], document_id))
    db.commit()
    flash("Flashcard deck mastery reset for this study note.", "success")
    return redirect(url_for("flashcards_view", document_id=document_id))


# 2. CONCEPT GRAPH ROUTE
@app.route("/documents/<int:document_id>/concept-graph")
@login_required()
def concept_graph_view(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    graph_data = extract_concept_graph(document["content"] or "", doc_title=document["title"])
    log_activity("concept_graph_view", document_id=document_id, detail=document["title"])
    return render_template(
        "concept_graph.html",
        document=document,
        graph_data=graph_data,
        graph_json=json.dumps(graph_data),
    )


# 3. 1-PAGE EXAM CHEAT SHEET ROUTE
@app.route("/documents/<int:document_id>/cheat-sheet")
@login_required()
def cheat_sheet_view(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    cheat_sheet_data = generate_cheat_sheet_data(document["content"] or "")
    summary = stored_summary(document_id)
    log_activity("cheat_sheet_view", document_id=document_id, detail=document["title"])
    return render_template(
        "cheat_sheet.html",
        document=document,
        cs=cheat_sheet_data,
        summary=summary,
    )


# 4. SOCRATIC VIVA VOCE SIMULATOR ROUTE
@app.route("/documents/<int:document_id>/viva")
@login_required()
def viva_view(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    questions = generate_viva_questions(document["content"] or "")
    db = get_db()
    history = db.execute(
        "SELECT * FROM viva_history WHERE user_id = ? AND document_id = ? ORDER BY created_at DESC LIMIT 10",
        (g.user["id"], document_id)
    ).fetchall()
    log_activity("viva_session_start", document_id=document_id, detail=document["title"])
    return render_template(
        "viva.html",
        document=document,
        questions=questions,
        history=history,
    )


@app.post("/documents/<int:document_id>/viva/evaluate")
@login_required()
def viva_evaluate(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    question_id = int(request.form.get("question_id", 1))
    student_answer = request.form.get("student_answer", "").strip()
    questions = generate_viva_questions(document["content"] or "")
    q_obj = next((q for q in questions if q["id"] == question_id), questions[0] if questions else None)
    if not q_obj:
        return jsonify({"error": "Question not found"}), 404
        
    eval_res = evaluate_viva_answer(q_obj["expected_keywords"], q_obj["model_answer"], student_answer)
    db = get_db()
    db.execute(
        "INSERT INTO viva_history (user_id, document_id, question, student_answer, score, feedback, created_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (g.user["id"], document_id, q_obj["question"], student_answer, eval_res["score"], eval_res["feedback"], now_iso())
    )
    db.commit()
    log_activity("viva_evaluation", document_id=document_id, detail=f"Score: {eval_res['score']}/10")
    return jsonify({
        "success": True,
        "evaluation": eval_res,
        "model_answer": q_obj["model_answer"],
    })


@app.post("/documents/<int:document_id>/viva/history/clear")
@login_required()
def viva_clear_history(document_id):
    document = get_document(document_id)
    if not visible_document(document):
        abort(403)
    db = get_db()
    db.execute("DELETE FROM viva_history WHERE user_id = ? AND document_id = ?", (g.user["id"], document_id))
    db.commit()
    flash("Viva practice history cleared.", "success")
    return redirect(url_for("viva_view", document_id=document_id))


@app.errorhandler(413)
def file_too_large(_error):
    flash("Files must be smaller than 10 MB.", "error")
    return redirect(request.referrer or url_for("dashboard"))


with app.app_context():
    init_db()

# Start scheduler for both direct run and gunicorn
start_scheduler()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
